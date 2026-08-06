"""Workspace-aware validation, recovery, generation and processing history."""
from __future__ import annotations

from copy import deepcopy
import logging
import os
from pathlib import Path
import shutil
from typing import Any, Dict, List, Mapping, Optional
import uuid

from fastapi import APIRouter, Request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from src.data_loader import DataLoader
from src.exporter import export_to_excel
from src.operator_issues import (
    attach_operator_issues,
    group_operator_issues,
    normalize_operator_issue,
)
from src.schedule_analyzer import ScheduleLayout
from src.schedule_period import resolve_schedule_calendar
from src.teacher_assignment import resolve_teacher_collisions
from src.transformer import transform_to_teacher_grid
from src.weekly_exporter import generate_weekly_semester_schedule
from src.workspace_domain import default_semester_settings
from web.backend.analysis_api import analyze_files, request_uploads
from web.backend.app_context import ApplicationContext
from web.backend.auth import actor_from_request
from web.backend.errors import ApplicationError
from web.backend.schemas import (
    FileUploadDetail,
    GenerateFileSpec,
    GenerateScheduleRequest,
    ScheduleUploadResponse,
    ValidateLayoutRequest,
    ValidateLayoutResponse,
)

logger = logging.getLogger(__name__)
UNASSIGNED_TEACHER = "Не назначен"


def _model_dict(model: Any) -> Dict[str, Any]:
    if hasattr(model, "model_dump"):
        return model.model_dump()
    return model.dict()


def _archive_source(
    context: ApplicationContext,
    run_id: str,
    file_id: str,
    source: Path,
) -> str:
    archive_root = (context.paths.data_dir / "history").resolve()
    destination_dir = archive_root / run_id / "sources"
    destination_dir.mkdir(parents=True, exist_ok=True)
    extension = source.suffix.lower() if source.suffix.lower() in {".xlsx", ".xlsm"} else ".xlsx"
    destination = destination_dir / f"{file_id}{extension}"
    if not destination.exists():
        try:
            os.link(source, destination)
        except OSError:
            shutil.copy2(source, destination)
    return destination.relative_to(context.paths.data_dir.resolve()).as_posix()


def _resolve_archive(context: ApplicationContext, relative: str) -> Path:
    root = context.paths.data_dir.resolve()
    path = (root / relative).resolve()
    history_root = (root / "history").resolve()
    if history_root not in path.parents or not path.is_file():
        raise ApplicationError(
            "Архивный исходный файл недоступен. Повторите загрузку вручную.",
            status_code=409,
            code="source_archive_missing",
        )
    return path


def _all_report_messages(report: Mapping[str, Any]) -> List[str]:
    result: List[str] = []
    for issue in report.get("issues") or []:
        if str(issue.get("severity")) == "info":
            continue
        text = str(issue.get("message") or "").strip()
        if text and text not in result:
            result.append(text)
    if result:
        return result
    for key in ("errors", "warnings"):
        for value in report.get(key) or []:
            text = str(value).strip()
            if text and "предварительная коллизия" not in text.casefold() and text not in result:
                result.append(text)
    return result


def _collect_corrections(
    reports: List[Dict[str, Any]],
    calendar_report: Mapping[str, Any],
    assignment_report: Mapping[str, Any] | None = None,
) -> List[Dict[str, Any]]:
    result: List[Dict[str, Any]] = []
    for report in reports:
        for item in report.get("auto_repairs") or []:
            result.append({"file": report.get("file"), **dict(item)})
    for item in calendar_report.get("corrections") or []:
        result.append(dict(item))
    for decision in (assignment_report or {}).get("decisions") or []:
        if decision.get("before") == decision.get("after"):
            continue
        result.append({
            "type": "teacher_reassignment",
            "field": "teacher",
            "subject": decision.get("subject"),
            "week": decision.get("week"),
            "day": decision.get("day"),
            "pair": decision.get("pair"),
            "groups": decision.get("groups") or [],
            "before": decision.get("before"),
            "after": decision.get("after"),
            "reason": decision.get("reason"),
            "blocking": False,
        })
    return result


def _teacher_overrides(values: Mapping[str, Any] | None) -> Dict[str, str]:
    raw = dict(values or {}).get("teacher_overrides") or {}
    if not isinstance(raw, Mapping):
        return {}
    return {
        str(subject).strip(): str(teacher).strip()
        for subject, teacher in raw.items()
        if str(subject).strip() and str(teacher).strip()
    }


def _record_teacher_overrides(
    loader: DataLoader,
    report: Dict[str, Any],
    requested: Mapping[str, str],
    applied: Mapping[str, str],
) -> None:
    report["teacher_overrides"] = dict(requested)
    for subject, requested_teacher in requested.items():
        resolved_teacher = applied.get(loader._subject_key(subject))
        if not resolved_teacher:
            continue
        report.setdefault("auto_repairs", []).append({
            "type": "teacher_override",
            "subject": subject,
            "before": "Не назначен",
            "after": resolved_teacher,
            "reason": f"Оператор назначил преподавателя для дисциплины «{subject}».",
            "blocking": False,
        })


def _assignment_for_group(
    assignment_report: Mapping[str, Any],
    group_name: str,
) -> Dict[str, Any]:
    decisions = [
        dict(item)
        for item in assignment_report.get("decisions") or []
        if group_name in (item.get("groups") or [])
    ]
    issues: List[Dict[str, Any]] = []
    for raw in assignment_report.get("issues") or []:
        items = [
            dict(item)
            for item in raw.get("items") or []
            if group_name in (item.get("groups") or [])
        ]
        if not items:
            continue
        issue = dict(raw)
        issue["items"] = items
        issues.append(issue)
    return {
        "algorithm": assignment_report.get("algorithm"),
        "decisions": decisions,
        "issues": issues,
        "reassigned_unit_count": sum(
            item.get("before") != item.get("after") for item in decisions
        ),
        "unresolved_conflict_count": sum(
            item.get("after") == UNASSIGNED_TEACHER for item in decisions
        ),
    }


def _report_issue_groups(
    reports: List[Mapping[str, Any]],
    assignment_report: Mapping[str, Any],
) -> List[Dict[str, Any]]:
    issues = [
        issue
        for report in reports
        for issue in report.get("issues") or []
    ]
    issues.extend(assignment_report.get("issues") or [])
    return group_operator_issues(issues)


def _actionable_messages(issue_groups: List[Mapping[str, Any]]) -> List[str]:
    result: List[str] = []
    for group in issue_groups:
        if not group.get("requires_action"):
            continue
        if group.get("category") == "calendar":
            continue
        for issue in group.get("items") or []:
            if issue.get("severity") == "info":
                continue
            message = str(issue.get("message") or "").strip()
            if message and message not in result:
                result.append(message)
    return result


def _export_recovery_workbook(
    output_path: Path,
    *,
    workspace_name: str,
    reports: List[Dict[str, Any]],
    calendar_report: Optional[Mapping[str, Any]] = None,
    assignment_report: Optional[Mapping[str, Any]] = None,
) -> None:
    """Create a useful result even when no lesson rows were recovered."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результат разбора"
    sheet.append(["Planner Solving", "Результат восстановления расписания"])
    sheet.append(["Рабочее пространство", workspace_name])
    sheet.append([
        "Статус",
        "Расписание пока не заполнено, но исходники сохранены и доступны для правки без повторной загрузки.",
    ])
    sheet.append([])
    sheet.append(["Файл", "Группа", "Найдено занятий", "Состояние", "Что можно сделать"])
    for report in reports:
        actions = "; ".join(
            str(item.get("label") or item.get("type") or "Уточнить")
            for item in report.get("actions") or []
        )
        sheet.append([
            report.get("file", ""),
            report.get("group", ""),
            int(report.get("lesson_count") or 0),
            report.get("status", "требует проверки"),
            actions or "Открыть файл в редакторе разметки",
        ])
    sheet.freeze_panes = "A6"
    for column, width in {"A": 34, "B": 20, "C": 18, "D": 24, "E": 65}.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[5]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    issues_sheet = workbook.create_sheet("Замечания и решения")
    issues_sheet.append([
        "Группа", "Файл", "Уровень", "Код", "Сообщение",
        "Решение по умолчанию", "Влияние", "Действие оператора",
    ])
    all_issues = [
        {"file": report.get("file"), **dict(item)}
        for report in reports
        for item in report.get("issues") or []
    ]
    all_issues.extend(
        {"file": "Все файлы", **normalize_operator_issue(item)}
        for item in (assignment_report or {}).get("issues") or []
    )
    for group in group_operator_issues(all_issues):
        for item in group.get("items") or []:
            issues_sheet.append([
                group.get("title"),
                item.get("file") or "Все файлы",
                item.get("severity", "attention"),
                item.get("code", "parser"),
                item.get("message", ""),
                item.get("default_decision", "Принято автоматически"),
                item.get("impact", "Формирование не блокируется"),
                (item.get("action") or {}).get("label") or "Необязательно",
            ])
    for item in (calendar_report or {}).get("issues") or []:
        issues_sheet.append([
            "Даты и календарь",
            "Все файлы",
            item.get("severity", "info"),
            item.get("code", "calendar"),
            item.get("message", ""),
            "Использовать фактическую последовательность дат",
            "Календарь записывается в отчёт без блокировки",
            (item.get("action") or {}).get("label") or "Необязательно",
        ])
    issues_sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFGH", (24, 28, 18, 30, 72, 54, 54, 40)):
        issues_sheet.column_dimensions[column].width = width
    for cell in issues_sheet[1]:
        cell.font = Font(bold=True)
    for row in issues_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(output_path)


def build_schedule_router(context: ApplicationContext) -> APIRouter:
    router = APIRouter(tags=["schedule"])
    repository = context.workspace_repository

    def workspace(workspace_id: str | None = None) -> Dict[str, Any]:
        return repository.get_workspace(workspace_id)

    def teacher_file(session_id: str, selected_workspace: Dict[str, Any]) -> Path:
        return context.sessions.write_snapshot(
            session_id,
            f"teachers-{selected_workspace['id']}.json",
            selected_workspace.get("teachers", []),
        )

    @router.post("/api/analysis/{session_id}/validate", response_model=ValidateLayoutResponse)
    def validate_layout(
        session_id: str,
        payload: ValidateLayoutRequest,
    ) -> ValidateLayoutResponse:
        manifest = context.sessions.load_manifest(session_id)
        item = context.sessions.manifest_file(manifest, payload.file_id)
        selected_workspace = workspace(payload.workspace_id)
        loader = DataLoader(str(teacher_file(session_id, selected_workspace)))
        requested_overrides = _teacher_overrides(payload.period_overrides)
        applied_overrides = loader.set_teacher_overrides(requested_overrides)
        lessons = loader.load_group_schedule(
            str(context.sessions.stored_path(session_id, item)),
            group_name=payload.group_name,
            layout=ScheduleLayout.from_dict(payload.layout),
            period_overrides=payload.period_overrides,
        )
        report = dict(loader.last_report)
        for warning in loader.warnings:
            if warning not in report.setdefault("warnings", []):
                report["warnings"].append(warning)
        _record_teacher_overrides(loader, report, requested_overrides, applied_overrides)
        report["used_lesson_count"] = len(lessons)
        attach_operator_issues(report)
        status = "success" if lessons and not _all_report_messages(report) else "warning"
        return ValidateLayoutResponse(status=status, report=report)

    def generate_from_session(
        session_id: str,
        payload: GenerateScheduleRequest,
        *,
        actor: Mapping[str, Any],
    ) -> ScheduleUploadResponse:
        manifest = context.sessions.load_manifest(session_id)
        specs = [spec for spec in payload.files if spec.enabled]
        if not specs:
            return ScheduleUploadResponse(
                status="warning",
                message="Не выбран ни один файл. Выберите файл или вернитесь к загрузке.",
            )

        selected_workspace = workspace(payload.workspace_id)
        run_id = repository.start_processing_run(
            workspace_id=selected_workspace["id"],
            session_id=session_id,
            actor=actor,
            source_count=len(specs),
        )
        teachers = list(selected_workspace.get("teachers", []))
        loader = DataLoader(str(teacher_file(session_id, selected_workspace)))
        lessons_all: List[Any] = []
        details: List[FileUploadDetail] = []
        reports: List[Dict[str, Any]] = []
        selected_templates: List[Dict[str, Any]] = []
        pending_files: List[Dict[str, Any]] = []
        global_overrides: Dict[str, str] = {}

        for spec in specs:
            item = context.sessions.manifest_file(manifest, spec.file_id)
            stored_path = context.sessions.stored_path(session_id, item)
            selected_match = dict((item.get("template_match") or {}).get("selected") or {})
            if selected_match.get("template_id"):
                selected_templates.append(selected_match)
            try:
                requested_overrides = _teacher_overrides(spec.period_overrides)
                applied_overrides = loader.set_teacher_overrides(requested_overrides)
                global_overrides.update(applied_overrides)
                warning_start = len(loader.warnings)
                lessons = loader.load_group_schedule(
                    str(stored_path),
                    group_name=spec.group_name,
                    layout=ScheduleLayout.from_dict(spec.layout),
                    period_overrides=spec.period_overrides,
                )
                report = dict(loader.last_report)
                report["group"] = spec.group_name
                for warning in loader.warnings[warning_start:]:
                    if warning not in report.setdefault("warnings", []):
                        report["warnings"].append(warning)
                _record_teacher_overrides(loader, report, requested_overrides, applied_overrides)
                report["period_overrides"] = dict(spec.period_overrides)
                report["source_archive"] = _archive_source(
                    context, run_id, spec.file_id, stored_path
                )
                reports.append(report)
                lessons_all.extend(lessons)
                pending_files.append({
                    "file_id": spec.file_id,
                    "filename": item["filename"],
                    "file_path": stored_path,
                    "group_name": spec.group_name,
                    "layout": report.get("layout_used") or spec.layout,
                    "report": report,
                    "template_match": selected_match,
                })
                details.append(FileUploadDetail(
                    file_id=spec.file_id,
                    filename=item["filename"],
                    status="success" if lessons else "warning",
                    message=(
                        f"Использовано занятий: {len(lessons)}."
                        if lessons
                        else "Занятия пока не извлечены; файл сохранён для правки на текущем экране."
                    ),
                    lesson_count=len(lessons),
                    used=bool(lessons),
                ))
            except Exception as exc:
                logger.exception(
                    "Cannot parse %s in analysis session %s",
                    item.get("filename", spec.file_id),
                    session_id,
                )
                error_report = {
                    "file": item.get("filename", spec.file_id),
                    "group": spec.group_name,
                    "errors": [f"Файл пропущен из-за технической ошибки: {exc}"],
                    "warnings": [],
                    "actions": [{
                        "type": "edit_layout",
                        "label": "Повторить разбор этого файла",
                        "blocking": False,
                    }],
                    "status": "skipped",
                    "blocking": False,
                    "period_overrides": dict(spec.period_overrides),
                }
                try:
                    error_report["source_archive"] = _archive_source(
                        context, run_id, spec.file_id, stored_path
                    )
                except OSError:
                    pass
                reports.append(error_report)
                pending_files.append({
                    "file_id": spec.file_id,
                    "filename": item["filename"],
                    "file_path": stored_path,
                    "group_name": spec.group_name,
                    "layout": spec.layout,
                    "report": error_report,
                    "template_match": selected_match,
                })
                details.append(FileUploadDetail(
                    file_id=spec.file_id,
                    filename=item["filename"],
                    status="warning",
                    message="Этот файл пропущен, остальные файлы продолжают обрабатываться.",
                    lesson_count=0,
                    used=False,
                    warning_count=1,
                    action_count=1,
                ))

        assignment_report = resolve_teacher_collisions(
            lessons_all,
            teachers,
            candidate_catalog=loader.teacher_candidate_catalog,
            manual_overrides=global_overrides,
        )

        for report, detail in zip(reports, details):
            group_name = str(report.get("group") or "")
            report["warnings"] = [
                warning
                for warning in report.get("warnings") or []
                if "предварительная коллизия" not in str(warning).casefold()
            ]
            assignment = _assignment_for_group(assignment_report, group_name)
            report["teacher_assignment"] = assignment
            report["assignment_issues"] = assignment["issues"]
            attach_operator_issues(report)
            messages = _all_report_messages(report)
            detail.warning_count = len(messages)
            detail.action_count = sum(
                bool(issue.get("actions")) for issue in report.get("issues") or []
            )
            if messages:
                detail.status = "warning"
                detail.message = (
                    f"Использовано занятий: {detail.lesson_count or 0}. "
                    f"Требуют внимания: {len(messages)}."
                )

        for pending in pending_files:
            repository.record_processing_file(
                run_id,
                file_id=pending["file_id"],
                filename=pending["filename"],
                file_path=pending["file_path"],
                group_name=pending["group_name"],
                layout=pending["layout"],
                report=pending["report"],
                template_match=pending["template_match"],
            )

        settings = selected_workspace.get("settings") or {}
        defaults = default_semester_settings()
        start_date = str(settings.get("schedule_start_date") or defaults["schedule_start_date"])
        end_date = str(settings.get("schedule_end_date") or defaults["schedule_end_date"])
        calendar = resolve_schedule_calendar(
            lessons_all,
            start_date_str=start_date,
            end_date_str=end_date,
            period_reports=reports,
            overrides=payload.calendar_overrides,
        )

        issue_groups = _report_issue_groups(reports, assignment_report)
        warnings = _actionable_messages(issue_groups)

        output_id = str(uuid.uuid4())
        filename = f"schedule_{output_id}.xlsx"
        output_path = context.paths.output_dir / filename
        weekly_filename: Optional[str] = None
        artifacts: List[Dict[str, Any]] = []

        if lessons_all:
            export_teachers = list(teachers)
            unknown_found = False
            for lesson in lessons_all:
                if str(lesson.teacher).strip().casefold() in {
                    "", "unknown", "none", UNASSIGNED_TEACHER.casefold()
                }:
                    lesson.teacher = UNASSIGNED_TEACHER
                    unknown_found = True
            if unknown_found and not any(
                item.get("short_name") == UNASSIGNED_TEACHER
                for item in export_teachers
            ):
                export_teachers.append({
                    "id": -1,
                    "short_name": UNASSIGNED_TEACHER,
                    "full_name": UNASSIGNED_TEACHER,
                    "position": "Требует уточнения оператором",
                    "rank": "—",
                    "academic_degree": "",
                })
            transformed = transform_to_teacher_grid(
                lessons_all,
                export_teachers,
                start_date_str=start_date,
                end_date_str=end_date,
                period_reports=reports,
                resolved_calendar=calendar,
            )
            transformed["teacher_assignment"] = assignment_report
            export_to_excel(transformed, export_teachers, str(output_path))
            artifacts.append({"kind": "schedule", "path": output_path})

            if context.paths.weekly_template.exists():
                weekly_filename = f"weekly_schedule_{output_id}.xlsx"
                weekly_path = context.paths.output_dir / weekly_filename
                try:
                    generate_weekly_semester_schedule(
                        teachers_config=export_teachers,
                        lessons=lessons_all,
                        template_path=str(context.paths.weekly_template),
                        output_path=str(weekly_path),
                        start_date_str=start_date,
                        end_date_str=end_date,
                        period_reports=reports,
                        resolved_calendar=calendar,
                    )
                    artifacts.append({"kind": "weekly", "path": weekly_path})
                except Exception as exc:
                    logger.exception("Cannot generate weekly schedule")
                    weekly_filename = None
                    warnings.append(f"Недельный файл не сформирован: {exc}")
            else:
                warnings.append(
                    "Недельный файл не сформирован: отсутствует шаблон obrazec/Недельное.xlsx."
                )
        else:
            filename = f"schedule_recovery_{output_id}.xlsx"
            output_path = context.paths.output_dir / filename
            _export_recovery_workbook(
                output_path,
                workspace_name=selected_workspace["name"],
                reports=reports,
                calendar_report=calendar.report,
                assignment_report=assignment_report,
            )
            artifacts.append({"kind": "recovery", "path": output_path})
            warnings.append(
                "Занятия не извлечены; сформирован диагностический файл. "
                "Исходники остаются в текущем сеансе для ручной коррекции."
            )

        corrections = _collect_corrections(
            reports, calendar.report, assignment_report
        )
        has_attention = bool(
            warnings
            or assignment_report.get("unresolved_conflict_count")
            or any(not item.used for item in details)
        )
        status = "warning" if has_attention else "success"
        error_count = sum(
            group.get("severity") in {"problem", "technical"}
            for group in issue_groups
        )
        run_report = {
            "reports": reports,
            "details": [_model_dict(item) for item in details],
            "period": calendar.report,
            "calendar_overrides": dict(payload.calendar_overrides),
            "corrections": corrections,
            "teacher_assignment": assignment_report,
            "issue_groups": issue_groups,
            "session_id": session_id,
        }
        repository.finish_processing_run(
            run_id,
            status=status,
            lesson_count=len(lessons_all),
            warning_count=len(warnings),
            error_count=error_count,
            selected_templates=selected_templates,
            report=run_report,
            artifacts=artifacts,
            actor=actor,
        )
        return ScheduleUploadResponse(
            run_id=run_id,
            filename=filename,
            weekly_filename=weekly_filename,
            status=status,
            message=(
                f"Расписание пространства «{selected_workspace['name']}» сформировано. "
                f"Использовано занятий: {len(lessons_all)}. "
                f"Автоматически устранено коллизий: "
                f"{assignment_report.get('resolved_conflict_count', 0)}."
                if lessons_all
                else (
                    f"Разбор пространства «{selected_workspace['name']}» завершён. "
                    "Создан файл с результатами диагностики и действиями оператора."
                )
            ),
            details=details,
            warnings=warnings,
            reports=reports,
            corrections=corrections,
            teacher_assignment=assignment_report,
            issue_groups=issue_groups,
        )

    @router.post("/api/analysis/{session_id}/generate", response_model=ScheduleUploadResponse)
    def generate_schedule(
        session_id: str,
        payload: GenerateScheduleRequest,
        request: Request,
    ) -> ScheduleUploadResponse:
        return generate_from_session(
            session_id, payload, actor=actor_from_request(request)
        )

    @router.post(
        "/api/workspaces/{workspace_id}/runs/{run_id}/repeat",
        response_model=ScheduleUploadResponse,
    )
    def repeat_processing_run(
        workspace_id: str,
        run_id: str,
        request: Request,
    ) -> ScheduleUploadResponse:
        actor = actor_from_request(request)
        previous = repository.get_processing_run(run_id)
        if previous["workspace_id"] != workspace_id:
            raise ApplicationError(
                "Запуск относится к другому пространству.", status_code=404
            )
        session_id, session_dir = context.sessions.create()
        manifest_files: List[Dict[str, Any]] = []
        specs: List[GenerateFileSpec] = []
        for source in previous.get("files", []):
            archived = _resolve_archive(
                context,
                str((source.get("report") or {}).get("source_archive") or ""),
            )
            file_id = str(uuid.uuid4())
            destination = session_dir / f"{file_id}{archived.suffix.lower()}"
            try:
                os.link(archived, destination)
            except OSError:
                shutil.copy2(archived, destination)
            manifest_files.append({
                "file_id": file_id,
                "filename": source["filename"],
                "group_name": source.get("group_name") or Path(source["filename"]).stem,
                "stored_name": destination.name,
                "status": "success",
                "message": "Восстановлено из истории обработки.",
                "analysis": None,
                "template_match": {
                    "selected": {
                        "template_id": source.get("template_id"),
                        "revision_no": source.get("template_revision"),
                        "score": source.get("match_score"),
                    }
                },
            })
            source_report = source.get("report") or {}
            specs.append(GenerateFileSpec(
                file_id=file_id,
                group_name=source.get("group_name") or Path(source["filename"]).stem,
                layout=source.get("layout") or source_report.get("layout_used") or {},
                enabled=True,
                period_overrides=source_report.get("period_overrides") or {},
            ))
        context.sessions.save_manifest(
            session_id,
            {
                "session_id": session_id,
                "created_at": previous.get("created_at"),
                "files": manifest_files,
            },
        )
        repository.audit(
            actor=actor,
            action="processing.repeat",
            entity_type="processing_run",
            entity_id=run_id,
            workspace_id=workspace_id,
            summary="Повтор обработки с прежними исходниками, правками и разметкой",
        )
        previous_report = previous.get("report") or {}
        return generate_from_session(
            session_id,
            GenerateScheduleRequest(
                files=specs,
                workspace_id=workspace_id,
                calendar_overrides=previous_report.get("calendar_overrides") or {},
                allow_partial=True,
            ),
            actor=actor,
        )

    @router.post("/api/upload", response_model=ScheduleUploadResponse)
    async def upload_compatibility(http_request: Request) -> ScheduleUploadResponse:
        analysis = await analyze_files(
            context, await request_uploads(http_request)
        )
        specs = [{
            "file_id": item.file_id,
            "group_name": item.group_name,
            "layout": item.analysis["layout"],
            "enabled": True,
            "period_overrides": {},
        } for item in analysis.files if item.analysis]
        generation_request = GenerateScheduleRequest(
            files=specs,
            workspace_id=repository.default_workspace_id(),
            allow_partial=True,
        )
        return generate_from_session(
            analysis.session_id,
            generation_request,
            actor=actor_from_request(http_request),
        )

    return router
