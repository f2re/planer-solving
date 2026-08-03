"""Excel schedule parser driven by an operator-confirmed exact layout."""
from __future__ import annotations

import logging
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Union

from openpyxl import load_workbook

from .schedule_analyzer import LESSON_CODE_RE, ScheduleAnalyzer, ScheduleLayout, WorksheetMatrix, normalize_text
from .schedule_parser_geometry import ScheduleParserGeometry
from .teacher_resolver import TeacherResolver

logger = logging.getLogger(__name__)


@dataclass
class Lesson:
    group: str
    subject: str
    lesson_type_code: str
    room: str
    week: int
    day_of_week: str
    pair_num: int
    teacher: str
    date_day: int
    month: str
    semester_info: str = ""
    year_info: str = ""


class DataLoader(TeacherResolver, ScheduleParserGeometry):
    def load_group_schedule(
        self,
        file_path: str,
        group_name: Optional[str] = None,
        layout: Optional[Union[ScheduleLayout, Dict[str, Any]]] = None,
    ) -> List[Lesson]:
        path = Path(file_path)
        group_name = normalize_text(group_name) or path.stem
        report: Dict[str, Any] = {
            "file": path.name,
            "group": group_name,
            "lesson_count": 0,
            "mapped_lessons": 0,
            "unknown_teacher_lessons": 0,
            "unknown_subjects": [],
            "legend_entries": 0,
            "empty_week_columns": [],
            "skipped_non_lesson_cells": 0,
            "warnings": [],
            "errors": [],
            "samples": [],
        }
        self.last_report = report
        try:
            workbook = load_workbook(path, data_only=True, read_only=False)
        except Exception as exc:
            report["errors"].append(f"Не удалось открыть Excel-файл: {exc}")
            return []
        if layout is None:
            analysis = ScheduleAnalyzer().analyze(str(path))
            selected = analysis.layout
            report["warnings"].extend(item.message for item in analysis.diagnostics if item.severity != "info")
        else:
            selected = layout if isinstance(layout, ScheduleLayout) else ScheduleLayout.from_dict(layout)
        if selected.sheet_name not in workbook.sheetnames:
            report["errors"].append(f"Лист «{selected.sheet_name}» отсутствует в книге.")
            return []
        sheet = workbook[selected.sheet_name]
        matrix = WorksheetMatrix(sheet)
        validation = selected.validate(sheet)
        report["warnings"].extend(item.message for item in validation if item.severity == "warning")
        report["errors"].extend(item.message for item in validation if item.severity == "error")
        if report["errors"]:
            return []
        semester, year = self._metadata(matrix)
        legend = self._legend(matrix, selected, report)
        weeks = self._weeks(matrix, selected, report)
        if not weeks:
            return []
        months, current_month = {}, "Unknown"
        for _, header_col, data_col in weeks:
            if selected.months_row:
                current_month = self._month(matrix.value(selected.months_row, header_col), current_month)
            months[data_col] = current_month
        lessons: List[Lesson] = []
        unknown: Set[str] = set()
        grid_end = selected.grid_end_row or sheet.max_row
        day_rows = selected.resolved_day_rows()
        pair_offsets = selected.resolved_pair_offsets()
        day_names = selected.day_names[: len(day_rows)]
        for day_name, day_start in zip(day_names, day_rows):
            if day_start > grid_end or day_start > sheet.max_row:
                continue
            date_row = day_start + selected.date_row_offset
            for pair_index, pair_offset in enumerate(pair_offsets):
                pair = pair_index + 1
                base = day_start + pair_offset
                if base > grid_end or base > sheet.max_row:
                    continue
                for week, _, data_col in weeks:
                    schedule_teachers: List[str] = []
                    if selected.cell_mode == "combined_cell":
                        code, subject, room, schedule_teachers = self._combined(matrix.value(base + selected.code_row_offset, data_col))
                    else:
                        code = normalize_text(matrix.value(base + selected.code_row_offset, data_col))
                        subject = normalize_text(matrix.value(base + selected.subject_row_offset, data_col))
                        room = normalize_text(matrix.value(base + selected.room_row_offset, data_col))
                        if selected.teacher_row_offset is not None:
                            schedule_teachers = self._extract_teachers(matrix.value(base + selected.teacher_row_offset, data_col))
                    if not code and not subject:
                        continue
                    if not LESSON_CODE_RE.search(code):
                        report["skipped_non_lesson_cells"] += 1
                        continue
                    if not subject:
                        parsed_code, parsed_subject, parsed_room, parsed_teachers = self._combined(code)
                        if parsed_subject:
                            code, subject, room = parsed_code, parsed_subject, room or parsed_room
                            schedule_teachers.extend(parsed_teachers)
                    if not subject or self._subject_key(subject) in self.NON_LESSON_SUBJECTS:
                        report["skipped_non_lesson_cells"] += 1
                        continue
                    lesson_type = self._lesson_type(code)
                    role = self._teacher_role(code)
                    entry = legend.get(self._subject_key(subject), {})
                    legend_teachers = entry.get(role, [])
                    if selected.teacher_source == "schedule":
                        candidates = schedule_teachers
                    elif selected.teacher_source == "both":
                        candidates = list(dict.fromkeys([*schedule_teachers, *legend_teachers]))
                    else:
                        candidates = legend_teachers
                    teacher = self._assign_teacher(week, day_name, pair, subject, lesson_type, candidates)
                    if teacher == "Unknown":
                        report["unknown_teacher_lessons"] += 1
                        unknown.add(subject)
                    else:
                        report["mapped_lessons"] += 1
                    lesson = Lesson(
                        group_name,
                        subject,
                        code,
                        room,
                        week,
                        day_name,
                        pair,
                        teacher,
                        self._day_number(matrix.value(date_row, data_col)) if 1 <= date_row <= sheet.max_row else 0,
                        months.get(data_col, "Unknown"),
                        semester,
                        year,
                    )
                    lessons.append(lesson)
                    if len(report["samples"]) < 12:
                        report["samples"].append(asdict(lesson))
        report["lesson_count"] = len(lessons)
        report["unknown_subjects"] = sorted(unknown)
        if not lessons:
            report["errors"].append("По заданной разметке не найдено ни одного занятия.")
        elif report["unknown_teacher_lessons"]:
            report["warnings"].append(f"Для {report['unknown_teacher_lessons']} занятий не удалось определить преподавателя.")
        return lessons
