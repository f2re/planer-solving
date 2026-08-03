from __future__ import annotations

from copy import copy
from datetime import datetime, timedelta
from typing import Any, Dict, List, Mapping, Optional, Sequence

import openpyxl
from openpyxl.styles import Alignment, Font

from src.data_loader import Lesson
from src.schedule_period import (
    DAY_NAMES,
    MONTH_GENITIVE,
    ResolvedScheduleCalendar,
    resolve_schedule_calendar,
)


def get_week_dates(start_date: datetime, week_num: int) -> List[datetime]:
    """Return Monday-Saturday for a schedule week, including week 0."""

    monday = start_date - timedelta(days=start_date.weekday())
    target_monday = monday + timedelta(weeks=week_num - 1)
    return [target_monday + timedelta(days=index) for index in range(6)]


def copy_cell(source_cell: Any, target_cell: Any) -> None:
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)
    target_cell.value = source_cell.value


def _configured_weeks(start_date: datetime, end_date: datetime) -> List[int]:
    current_date = start_date - timedelta(days=start_date.weekday())
    weeks: List[int] = []
    week_num = 1
    while current_date <= end_date:
        weeks.append(week_num)
        week_num += 1
        current_date += timedelta(weeks=1)
    return weeks


def generate_weekly_semester_schedule(
    teachers_config: List[Dict[str, Any]],
    lessons: List[Lesson],
    template_path: str,
    output_path: str,
    start_date_str: str,
    end_date_str: str,
    *,
    period_reports: Optional[Sequence[Mapping[str, Any]]] = None,
    resolved_calendar: Optional[ResolvedScheduleCalendar] = None,
) -> None:
    calendar = resolved_calendar or resolve_schedule_calendar(
        lessons,
        start_date_str=start_date_str,
        end_date_str=end_date_str,
        period_reports=period_reports,
    )

    template_wb = openpyxl.load_workbook(template_path)
    template_ws = template_wb.active
    output_wb = openpyxl.Workbook()
    output_wb.remove(output_wb.active)

    weeks = list(calendar.weeks)
    day_map = {"Пн": 0, "Вт": 1, "Ср": 2, "Чт": 3, "Пт": 4, "Сб": 5}
    schedule_grid: Dict[tuple[int, str, int, int], List[Lesson]] = {}
    for lesson in lessons:
        day_index = day_map.get(lesson.day_of_week, -1)
        if day_index < 0 or lesson.week not in weeks:
            continue
        key = (lesson.week, lesson.teacher, day_index, lesson.pair_num)
        schedule_grid.setdefault(key, []).append(lesson)

    day_names_full = [
        "Понедельник",
        "Вторник",
        "Среда",
        "Четверг",
        "Пятница",
        "Суббота",
    ]
    pair_labels = {1: "1-2", 2: "3-4", 3: "5-6", 4: "7-8"}

    for week_num in weeks:
        worksheet = output_wb.create_sheet(title=f"Неделя {week_num}")
        for column in range(1, 11):
            letter = openpyxl.utils.get_column_letter(column)
            worksheet.column_dimensions[letter].width = template_ws.column_dimensions[letter].width

        for row in range(1, 13):
            for column in range(1, 11):
                copy_cell(
                    template_ws.cell(row=row, column=column),
                    worksheet.cell(row=row, column=column),
                )
            if row in template_ws.row_dimensions:
                worksheet.row_dimensions[row].height = template_ws.row_dimensions[row].height

        for day_index, day_name in enumerate(DAY_NAMES):
            day = calendar.date_for(week_num, day_name)
            if day is None:
                continue
            column = 5 + day_index
            worksheet.cell(row=11, column=column).value = (
                f"{day_names_full[day_index]}\n"
                f"({day.strftime('%d.%m')} · {MONTH_GENITIVE[day.month]})"
            )
            worksheet.cell(row=11, column=column).alignment = Alignment(
                wrapText=True,
                horizontal="center",
                vertical="center",
            )

        current_row = 13
        for teacher_index, teacher_info in enumerate(teachers_config):
            full_name_parts = str(teacher_info.get("full_name") or "").split()
            if len(full_name_parts) >= 3:
                name_with_initials = (
                    f"{full_name_parts[0]} "
                    f"{full_name_parts[1][0]}.{full_name_parts[2][0]}."
                )
            elif len(full_name_parts) == 2:
                name_with_initials = f"{full_name_parts[0]} {full_name_parts[1][0]}."
            else:
                name_with_initials = str(teacher_info.get("short_name") or "")

            teacher_short = str(teacher_info.get("short_name") or "")
            academic_rank = str(teacher_info.get("rank") or "-")
            for pair_index in range(1, 5):
                row = current_row + pair_index - 1
                for column in range(1, 11):
                    copy_cell(
                        template_ws.cell(row=14, column=column),
                        worksheet.cell(row=row, column=column),
                    )
                    worksheet.cell(row=row, column=column).value = None

                if pair_index == 1:
                    worksheet.cell(row=row, column=1).value = teacher_index + 1
                    worksheet.cell(row=row, column=2).value = academic_rank
                    worksheet.cell(row=row, column=3).value = name_with_initials
                worksheet.cell(row=row, column=4).value = pair_labels[pair_index]

                for day_index in range(6):
                    column = 5 + day_index
                    slot_lessons = schedule_grid.get(
                        (week_num, teacher_short, day_index, pair_index),
                        [],
                    )
                    if not slot_lessons:
                        continue
                    worksheet.cell(row=row, column=column).value = "\n".join(
                        f"{lesson.subject}, {lesson.group}, {lesson.room}"
                        for lesson in slot_lessons
                    )
                    worksheet.cell(row=row, column=column).alignment = Alignment(
                        wrapText=True,
                        horizontal="center",
                        vertical="center",
                    )
                    worksheet.cell(row=row, column=column).font = Font(
                        name="Arial",
                        size=7,
                    )

            worksheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row + 3,
                end_column=1,
            )
            worksheet.merge_cells(
                start_row=current_row,
                start_column=2,
                end_row=current_row + 3,
                end_column=2,
            )
            worksheet.merge_cells(
                start_row=current_row,
                start_column=3,
                end_row=current_row + 3,
                end_column=3,
            )
            current_row += 4

    output_wb.save(output_path)
