"""Template-independent discovery and preview of Excel schedule layouts."""
from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass
from datetime import date, datetime
import re
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .schedule_detection import ScheduleStructureDetector
from .schedule_layout import (
    DAY_NAMES, MONTH_PREFIXES, LESSON_CODE_RE, PAIR_LABEL_RE, DAY_PATTERNS,
    Diagnostic, ScheduleLayout, excel_display, is_legend_record, normalize_text, value_as_int,
)
from .worksheet_matrix import WorksheetMatrix


@dataclass
class AnalysisResult:
    sheet_names: List[str]
    selected_sheet: str
    max_row: int
    max_column: int
    confidence: float
    layout: ScheduleLayout
    diagnostics: List[Diagnostic]
    schedule_preview: Dict[str, int]
    legend_preview: Dict[str, int]

    def to_dict(self) -> Dict[str, Any]:
        return {
            "sheet_names": self.sheet_names,
            "selected_sheet": self.selected_sheet,
            "max_row": self.max_row,
            "max_column": self.max_column,
            "confidence": round(self.confidence, 3),
            "layout": self.layout.to_dict(),
            "diagnostics": [item.to_dict() for item in self.diagnostics],
            "schedule_preview": self.schedule_preview,
            "legend_preview": self.legend_preview,
        }


class ScheduleAnalyzer(ScheduleStructureDetector):
    def analyze(self, file_path: str, preferred_sheet: Optional[str] = None) -> AnalysisResult:
        workbook = load_workbook(file_path, data_only=True, read_only=False)
        if not workbook.sheetnames:
            raise ValueError("В книге нет листов.")
        selected = self._select_sheet(workbook, preferred_sheet)
        sheet = workbook[selected]
        matrix = WorksheetMatrix(sheet)
        diagnostics: List[Diagnostic] = []

        weeks_row, week_cols, week_numbers, week_score = self._detect_weeks(matrix)
        if weeks_row is None or not week_cols:
            weeks_row, week_cols, week_numbers = 1, [1], [1]
            diagnostics.append(Diagnostic("error", "weeks_not_found", "Не удалось надёжно определить строку учебных недель. Укажите её вручную."))
        month_row, month_score = self._detect_months(matrix, weeks_row, week_cols)
        day_rows = self._detect_day_rows(matrix, weeks_row, week_cols[0])
        day_gaps = [right[0] - left[0] for left, right in zip(day_rows, day_rows[1:])]
        day_block = self._median(day_gaps, 13)
        if day_rows:
            pair_offsets = self._detect_pair_offsets(
                matrix,
                day_rows[0][0],
                day_rows[1][0] if len(day_rows) > 1 else None,
                week_cols[0],
            )
            grid_start = day_rows[0][0]
        else:
            pair_offsets = [0, 3, 6, 9]
            grid_start = weeks_row + 3
            diagnostics.append(Diagnostic("warning", "day_labels_not_found", "Подписи дней недели не найдены; используется безопасная геометрия 13 строк и 4 пары."))
        pair_stride = self._median([b - a for a, b in zip(pair_offsets, pair_offsets[1:])], 3)

        sample_bases = []
        for day_row, _ in day_rows[:3] or [(grid_start, "Пн")]:
            sample_bases.extend(day_row + offset for offset in pair_offsets)
        structural_week_cols = self._week_candidate_columns(matrix, weeks_row, week_cols)
        code_scores, subject_scores, room_scores = self._row_kind_scores(
            matrix, sample_bases, structural_week_cols
        )
        code_offset = max(code_scores, key=code_scores.get, default=0)
        subject_candidates = {key: value for key, value in subject_scores.items() if key != code_offset}
        subject_offset = max(subject_candidates, key=subject_candidates.get, default=code_offset + 1)
        room_candidates = {key: value for key, value in room_scores.items() if key not in {code_offset, subject_offset}}
        room_offset = max(room_candidates, key=room_candidates.get, default=subject_offset + 1)

        week_data_cols = self._detect_week_data_columns(
            matrix,
            weeks_row,
            week_cols,
            day_rows,
            pair_offsets,
            code_offset,
            subject_offset,
        )

        date_scores = Counter()
        for offset in range(-3, 2):
            for day_row, _ in day_rows or [(grid_start, "Пн")]:
                row = day_row + offset
                if not 1 <= row <= sheet.max_row:
                    continue
                for col in week_data_cols:
                    value = matrix.value(row, col)
                    if isinstance(value, (date, datetime)):
                        date_scores[offset] += 3
                    elif re.fullmatch(r"\d{1,2}(?:[./-]\d{1,2})?", normalize_text(value)):
                        date_scores[offset] += 1
        date_offset = max(date_scores, key=date_scores.get, default=-1)

        legend = self._detect_legend(matrix, weeks_row + 10)
        if not legend.get("start_row"):
            diagnostics.append(Diagnostic("warning", "legend_not_found", "Блок дисциплин и преподавателей не найден. Укажите его вручную."))
        grid_end = (
            legend["start_row"] - 1
            if legend.get("start_row")
            else min(sheet.max_row, (day_rows[-1][0] if day_rows else grid_start) + max(pair_offsets) + max(room_offset, 2))
        )
        step = self._median([b - a for a, b in zip(week_cols, week_cols[1:])], 1)
        layout = ScheduleLayout(
            sheet_name=selected,
            weeks_row=weeks_row,
            first_week_col=week_cols[0],
            last_week_col=week_cols[-1],
            week_col_step=step,
            week_columns=list(week_cols),
            week_data_columns=list(week_data_cols),
            week_numbers=list(week_numbers),
            allow_week_zero=0 in week_numbers,
            months_row=month_row,
            grid_start_row=grid_start,
            grid_end_row=max(grid_start, grid_end),
            day_block_rows=day_block,
            day_start_rows=[row for row, _ in day_rows],
            pairs_per_day=len(pair_offsets),
            pair_row_stride=pair_stride,
            pair_row_offsets=list(pair_offsets),
            date_row_offset=date_offset,
            code_row_offset=code_offset,
            subject_row_offset=subject_offset,
            room_row_offset=room_offset,
            day_names=[day for _, day in day_rows] or list(DAY_NAMES),
            legend_start_row=legend.get("start_row"),
            legend_end_row=legend.get("end_row"),
            legend_data_start_offset=legend.get("data_start_offset", 3),
            legend_data_start_row=legend.get("data_start_row"),
            legend_code_col=legend.get("code_col"),
            legend_subject_col=legend.get("subject_col"),
            legend_lecturer_col=legend.get("lecturer_col"),
            legend_other_col=legend.get("other_col"),
            teacher_role_fallback="strict",
        )
        confidence = min(
            1.0,
            0.42 * week_score
            + 0.18 * month_score
            + 0.22 * min(1.0, len(day_rows) / 6)
            + 0.18 * legend.get("score", 0),
        )
        if confidence < 0.6:
            diagnostics.insert(0, Diagnostic("warning", "low_confidence", "Автоматическая разметка имеет низкую уверенность. Перед обработкой проверьте границы."))
        if len(day_rows) == 6:
            diagnostics.append(Diagnostic("info", "exact_day_rows", "Найдены точные строки всех шести учебных дней."))
        if week_cols:
            diagnostics.append(Diagnostic("info", "exact_week_columns", f"Найдено точных столбцов недель: {len(week_cols)}."))
        if 0 in week_numbers:
            diagnostics.append(Diagnostic("info", "week_zero", "В книге присутствует неделя 0; она сохранена в разметке."))

        schedule_preview = {
            "row_start": max(1, min(weeks_row, grid_start) - 2),
            "row_end": min(sheet.max_row, max(layout.grid_end_row or grid_start, grid_start) + 2),
            "col_start": max(1, week_cols[0] - 4),
            "col_end": min(sheet.max_column, week_cols[-1] + 2),
        }
        legend_start = legend.get("start_row") or min(sheet.max_row, (layout.grid_end_row or 1) + 1)
        legend_end = legend.get("end_row") or min(sheet.max_row, legend_start + 15)
        legend_cols = [
            value for value in (
                legend.get("code_col"), legend.get("subject_col"),
                legend.get("lecturer_col"), legend.get("other_col"),
            ) if value
        ]
        legend_preview = {
            "row_start": max(1, legend_start - 2),
            "row_end": min(sheet.max_row, legend_end + 2),
            "col_start": max(1, min(legend_cols or [1]) - 1),
            "col_end": min(sheet.max_column, max(legend_cols or [min(sheet.max_column, 12)]) + 2),
        }
        return AnalysisResult(
            workbook.sheetnames,
            selected,
            sheet.max_row,
            sheet.max_column,
            confidence,
            layout,
            diagnostics,
            schedule_preview,
            legend_preview,
        )

    def preview(
        self,
        file_path: str,
        sheet_name: str,
        row_start: int,
        row_end: int,
        col_start: int,
        col_end: int,
        max_rows: int = 140,
        max_columns: int = 60,
    ) -> Dict[str, Any]:
        workbook = load_workbook(file_path, data_only=True, read_only=False)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Лист «{sheet_name}» не найден.")
        sheet = workbook[sheet_name]
        matrix = WorksheetMatrix(sheet)
        row_start = max(1, min(row_start, sheet.max_row))
        row_end = max(row_start, min(row_end, sheet.max_row, row_start + max_rows - 1))
        col_start = max(1, min(col_start, sheet.max_column))
        col_end = max(col_start, min(col_end, sheet.max_column, col_start + max_columns - 1))
        columns = [
            {
                "index": col,
                "label": get_column_letter(col),
                "width": sheet.column_dimensions[get_column_letter(col)].width or 10,
            }
            for col in range(col_start, col_end + 1)
        ]
        rows = []
        for row in range(row_start, row_end + 1):
            cells = []
            for col in range(col_start, col_end + 1):
                cell = sheet.cell(row, col)
                anchor_row, anchor_col = matrix.merged_anchor(row, col)
                color = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type and cell.fill.fgColor.type == "rgb" else None
                cells.append({
                    "column": col,
                    "value": excel_display(matrix.value(row, col)),
                    "merged": matrix.is_merged(row, col),
                    "anchor_row": anchor_row,
                    "anchor_column": anchor_col,
                    "is_anchor": matrix.is_anchor(row, col),
                    "bold": bool(cell.font and cell.font.bold),
                    "fill": color,
                    "alignment": cell.alignment.horizontal if cell.alignment else None,
                })
            rows.append({"index": row, "height": sheet.row_dimensions[row].height or 15, "cells": cells})
        return {
            "sheet_name": sheet_name,
            "row_start": row_start,
            "row_end": row_end,
            "col_start": col_start,
            "col_end": col_end,
            "columns": columns,
            "rows": rows,
        }
