"""Editable schedule-layout model and parser primitives."""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import date, datetime
import math
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl.worksheet.worksheet import Worksheet

DAY_NAMES = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб"]
MONTH_PREFIXES = {
    "янв": "Январь", "фев": "Февраль", "мар": "Март", "апр": "Апрель",
    "май": "Май", "июн": "Июнь", "июл": "Июль", "авг": "Август",
    "сен": "Сентябрь", "окт": "Октябрь", "ноя": "Ноябрь", "дек": "Декабрь",
}
LESSON_CODE_RE = re.compile(
    r"^(?:экзамен|конс|зач|икс|пз|пр|лр|лт|зч|зо|экз|кр|кп|гз|пп|"
    r"гу|л|п|с|у|э)"
    r"(?:\s*[/\\.\-]|\b)",
    re.IGNORECASE,
)
PAIR_LABEL_RE = re.compile(r"^\s*(\d+)\s*[-–]\s*(\d+)\s*$")
DAY_PATTERNS = {
    "Пн": re.compile(r"^(?:пн|понедельник)\.?$", re.I),
    "Вт": re.compile(r"^(?:вт|вторник)\.?$", re.I),
    "Ср": re.compile(r"^(?:ср|среда)\.?$", re.I),
    "Чт": re.compile(r"^(?:чт|четверг)\.?$", re.I),
    "Пт": re.compile(r"^(?:пт|пятница)\.?$", re.I),
    "Сб": re.compile(r"^(?:сб|суббота)\.?$", re.I),
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(
        r"\s+",
        " ",
        str(value).replace("\xa0", " ").replace("ё", "е").replace("Ё", "Е"),
    ).strip()


def value_as_int(value: Any) -> Optional[int]:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return int(value)
    text = normalize_text(value)
    if re.fullmatch(r"-?\d{1,3}(?:\.0+)?", text):
        return int(float(text))
    return None


def excel_display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


LEGEND_CODE_RE = re.compile(r"[0-9A-Za-zА-Яа-яЁё./_-]{1,24}")


def is_legend_record(code: Any, subject: Any, lecturer: Any, other: Any) -> bool:
    """Return True only for a real discipline row, not the explanation below it."""

    code_text = normalize_text(code)
    if not LEGEND_CODE_RE.fullmatch(code_text):
        return False
    subject_text = normalize_text(subject)
    lecturer_text = normalize_text(lecturer)
    other_text = normalize_text(other)
    if lecturer_text or other_text:
        return True
    code_key = re.sub(r"[^0-9a-zа-я]+", "", code_text.casefold().replace("ё", "е"))
    subject_key = re.sub(r"[^0-9a-zа-я]+", "", subject_text.casefold().replace("ё", "е"))
    return bool(subject_key and len(subject_key) >= 2 and subject_key != code_key)


@dataclass
class Diagnostic:
    severity: str
    code: str
    message: str
    row: Optional[int] = None
    column: Optional[int] = None

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class ScheduleLayout:
    """Operator-editable coordinates; all rows and columns are 1-based."""

    sheet_name: str = ""
    weeks_row: int = 1
    first_week_col: int = 1
    last_week_col: int = 1
    week_col_step: int = 1
    week_data_col_offset: int = 0
    week_columns: List[int] = field(default_factory=list)
    week_data_columns: List[int] = field(default_factory=list)
    week_numbers: List[int] = field(default_factory=list)
    allow_week_zero: bool = True
    months_row: Optional[int] = None
    grid_start_row: int = 1
    grid_end_row: Optional[int] = None
    day_block_rows: int = 13
    day_start_rows: List[int] = field(default_factory=list)
    pairs_per_day: int = 4
    pair_row_stride: int = 3
    pair_row_offsets: List[int] = field(default_factory=list)
    date_row_offset: int = -1
    code_row_offset: int = 0
    subject_row_offset: int = 1
    room_row_offset: int = 2
    teacher_row_offset: Optional[int] = None
    day_names: List[str] = field(default_factory=lambda: list(DAY_NAMES))
    legend_start_row: Optional[int] = None
    legend_end_row: Optional[int] = None
    legend_data_start_offset: int = 3
    legend_data_start_row: Optional[int] = None
    legend_code_col: Optional[int] = 1
    legend_subject_col: Optional[int] = None
    legend_lecturer_col: Optional[int] = None
    legend_other_col: Optional[int] = None
    cell_mode: str = "row_layers"
    teacher_source: str = "legend"
    teacher_role_fallback: str = "strict"

    @classmethod
    def from_dict(cls, raw: Dict[str, Any]) -> "ScheduleLayout":
        allowed = cls.__dataclass_fields__
        data = {key: value for key, value in raw.items() if key in allowed}
        data["day_names"] = [str(item) for item in (data.get("day_names") or DAY_NAMES)]
        for field_name in (
            "week_columns", "week_data_columns", "week_numbers",
            "day_start_rows", "pair_row_offsets",
        ):
            values = data.get(field_name) or []
            data[field_name] = [
                int(parsed) for item in values
                if (parsed := value_as_int(item)) is not None
            ]
        return cls(**data)

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)

    def resolved_week_columns(self) -> List[int]:
        if self.week_columns:
            return list(dict.fromkeys(int(item) for item in self.week_columns))
        return list(range(self.first_week_col, self.last_week_col + 1, max(1, self.week_col_step)))

    def resolved_day_rows(self) -> List[int]:
        if self.day_start_rows:
            return list(self.day_start_rows)
        return [self.grid_start_row + index * self.day_block_rows for index in range(len(self.day_names))]

    def resolved_pair_offsets(self) -> List[int]:
        if self.pair_row_offsets:
            return list(self.pair_row_offsets)
        return [index * self.pair_row_stride for index in range(self.pairs_per_day)]

    def repaired(self, worksheet: Worksheet) -> Tuple["ScheduleLayout", List[Dict[str, Any]]]:
        """Return a safe layout and a transparent list of automatic corrections.

        Operator coordinates are treated as hints. Obvious inversions, invalid
        values and out-of-sheet coordinates are repaired instead of rejecting the
        file. Every change is returned to the UI and can still be edited in place.
        """

        fixed = ScheduleLayout.from_dict(self.to_dict())
        repairs: List[Dict[str, Any]] = []
        max_row = max(1, int(worksheet.max_row or 1))
        max_col = max(1, int(worksheet.max_column or 1))

        def change(field_name: str, value: Any, reason: str) -> None:
            before = getattr(fixed, field_name)
            if before == value:
                return
            setattr(fixed, field_name, value)
            repairs.append({
                "type": "layout_patch",
                "field": field_name,
                "before": before,
                "after": value,
                "reason": reason,
                "blocking": False,
            })

        defaults = {
            "weeks_row": 1,
            "first_week_col": 1,
            "last_week_col": max_col,
            "week_col_step": 1,
            "grid_start_row": 1,
            "day_block_rows": 13,
            "pairs_per_day": 4,
            "pair_row_stride": 3,
        }
        for field_name, default in defaults.items():
            value = value_as_int(getattr(fixed, field_name))
            if value is None or value < 1:
                change(field_name, default, "Недопустимое значение заменено безопасным.")

        change("weeks_row", min(max(1, fixed.weeks_row), max_row), "Строка недель приведена к границам листа.")
        change("grid_start_row", min(max(1, fixed.grid_start_row), max_row), "Начало сетки приведено к границам листа.")
        change("first_week_col", min(max(1, fixed.first_week_col), max_col), "Первый столбец приведён к границам листа.")
        change("last_week_col", min(max(1, fixed.last_week_col), max_col), "Последний столбец приведён к границам листа.")
        if fixed.first_week_col > fixed.last_week_col:
            left, right = fixed.last_week_col, fixed.first_week_col
            change("first_week_col", left, "Перепутанные границы недель переставлены местами.")
            change("last_week_col", right, "Перепутанные границы недель переставлены местами.")

        for field_name in ("months_row", "grid_end_row", "legend_start_row", "legend_end_row", "legend_data_start_row"):
            value = getattr(fixed, field_name)
            if value is not None:
                parsed = value_as_int(value)
                change(field_name, min(max(1, parsed or 1), max_row), "Строка приведена к границам листа.")
        for field_name in ("legend_code_col", "legend_subject_col", "legend_lecturer_col", "legend_other_col"):
            value = getattr(fixed, field_name)
            if value is not None:
                parsed = value_as_int(value)
                change(field_name, min(max(1, parsed or 1), max_col), "Столбец приведён к границам листа.")

        if fixed.grid_end_row is not None and fixed.grid_end_row < fixed.grid_start_row:
            change("grid_end_row", max_row, "Конец сетки автоматически перенесён ниже её начала.")
        if fixed.legend_start_row and fixed.legend_end_row and fixed.legend_end_row < fixed.legend_start_row:
            change("legend_end_row", max(fixed.legend_start_row, fixed.legend_end_row), "Границы блока дисциплин упорядочены.")
        if fixed.legend_start_row and fixed.legend_data_start_row and fixed.legend_data_start_row <= fixed.legend_start_row:
            change(
                "legend_data_start_row",
                min(max_row, fixed.legend_start_row + max(1, fixed.legend_data_start_offset)),
                "Начало данных блока перенесено ниже заголовка.",
            )

        columns = sorted({value for value in fixed.resolved_week_columns() if 1 <= value <= max_col})
        if not columns:
            columns = list(range(fixed.first_week_col, fixed.last_week_col + 1, max(1, fixed.week_col_step)))
        if not columns:
            columns = [fixed.first_week_col]
        change("week_columns", columns, "Точные столбцы недель восстановлены по границам листа.")
        change("first_week_col", columns[0], "Граница синхронизирована с точными столбцами недель.")
        change("last_week_col", columns[-1], "Граница синхронизирована с точными столбцами недель.")

        data_columns = [value for value in fixed.week_data_columns if 1 <= value <= max_col]
        if len(data_columns) != len(columns):
            data_columns = [
                min(max_col, max(1, column + int(fixed.week_data_col_offset or 0)))
                for column in columns
            ]
        change("week_data_columns", data_columns, "Столбцы данных выровнены со столбцами учебных недель.")

        minimum_week = 0 if fixed.allow_week_zero else 1
        week_numbers = [value for value in fixed.week_numbers if minimum_week <= value <= 60]
        if len(week_numbers) != len(columns):
            first = week_numbers[0] if week_numbers else minimum_week
            week_numbers = [first + index for index in range(len(columns))]
        change("week_numbers", week_numbers, "Номера недель выровнены с числом столбцов.")

        if not fixed.day_names:
            change("day_names", list(DAY_NAMES), "Восстановлен стандартный список дней недели.")
        day_rows = [value for value in fixed.day_start_rows if 1 <= value <= max_row]
        if len(day_rows) != len(fixed.day_names):
            day_rows = [
                min(max_row, fixed.grid_start_row + index * max(1, fixed.day_block_rows))
                for index in range(len(fixed.day_names))
            ]
        change("day_start_rows", day_rows, "Строки дней восстановлены по началу и высоте блока дня.")

        pair_offsets = [value for value in fixed.pair_row_offsets if isinstance(value, int) and value >= 0]
        if len(pair_offsets) != max(1, fixed.pairs_per_day):
            pair_offsets = [index * max(1, fixed.pair_row_stride) for index in range(max(1, fixed.pairs_per_day))]
        change("pair_row_offsets", pair_offsets, "Смещения пар восстановлены по шагу строк.")

        if fixed.cell_mode not in {"row_layers", "combined_cell"}:
            change("cell_mode", "row_layers", "Неизвестный режим заменён послойным чтением.")
        if fixed.teacher_source not in {"legend", "schedule", "both"}:
            change("teacher_source", "both", "Преподаватель будет искаться и в сетке, и в блоке дисциплин.")
        if fixed.teacher_role_fallback not in {"strict", "any"}:
            change("teacher_role_fallback", "any", "Разрешена безопасная подстановка роли преподавателя.")
        return fixed, repairs

    def validate(self, worksheet: Optional[Worksheet] = None) -> List[Diagnostic]:
        result: List[Diagnostic] = []
        required = {
            "строка недель": self.weeks_row,
            "первый столбец недели": self.first_week_col,
            "последний столбец недели": self.last_week_col,
            "шаг столбцов недель": self.week_col_step,
            "начало сетки": self.grid_start_row,
            "высота блока дня": self.day_block_rows,
            "число пар": self.pairs_per_day,
            "шаг строки пары": self.pair_row_stride,
        }
        optional = {
            "строка месяцев": self.months_row,
            "конец сетки": self.grid_end_row,
            "начало легенды": self.legend_start_row,
            "конец легенды": self.legend_end_row,
            "первая строка данных легенды": self.legend_data_start_row,
            "столбец обозначения": self.legend_code_col,
            "столбец дисциплины": self.legend_subject_col,
            "столбец лектора": self.legend_lecturer_col,
            "столбец других преподавателей": self.legend_other_col,
        }
        for label, value in required.items():
            if not isinstance(value, int) or value < 1:
                result.append(Diagnostic("warning", "invalid_coordinate", f"Поле «{label}» будет исправлено автоматически."))
        for label, value in optional.items():
            if value is not None and (not isinstance(value, int) or value < 1):
                result.append(Diagnostic("warning", "invalid_coordinate", f"Поле «{label}» будет исправлено автоматически."))
        if self.first_week_col > self.last_week_col:
            result.append(Diagnostic("warning", "invalid_week_range", "Границы недель будут переставлены местами."))
        if self.grid_end_row and self.grid_start_row > self.grid_end_row:
            result.append(Diagnostic("warning", "invalid_grid_range", "Границы сетки будут упорядочены автоматически."))
        if self.legend_start_row and self.legend_end_row and self.legend_start_row > self.legend_end_row:
            result.append(Diagnostic("warning", "invalid_legend_range", "Границы блока дисциплин будут упорядочены автоматически."))
        if self.legend_data_start_row and self.legend_start_row and self.legend_data_start_row <= self.legend_start_row:
            result.append(Diagnostic("warning", "invalid_legend_data", "Начало данных будет перенесено ниже заголовка."))
        if self.cell_mode not in {"row_layers", "combined_cell"}:
            result.append(Diagnostic("warning", "invalid_cell_mode", "Режим ячеек будет заменён безопасным вариантом."))
        if self.teacher_source not in {"legend", "schedule", "both"}:
            result.append(Diagnostic("warning", "invalid_teacher_source", "Источник преподавателя будет определён автоматически."))
        if self.teacher_role_fallback not in {"strict", "any"}:
            result.append(Diagnostic("warning", "invalid_teacher_fallback", "Режим подстановки будет исправлен автоматически."))
        if not self.day_names:
            result.append(Diagnostic("warning", "missing_days", "Будет использован стандартный список дней недели."))
        for name, values in (
            ("точные столбцы недель", self.week_columns),
            ("столбцы данных недель", self.week_data_columns),
            ("строки дней", self.day_start_rows),
        ):
            if any(not isinstance(value, int) or value < 1 for value in values):
                result.append(Diagnostic("warning", "invalid_coordinate_list", f"Поле «{name}» будет очищено и восстановлено."))
        if any(not isinstance(value, int) or value < 0 for value in self.pair_row_offsets):
            result.append(Diagnostic("warning", "invalid_pair_offsets", "Смещения пар будут восстановлены по шагу строк."))
        minimum_week = 0 if self.allow_week_zero else 1
        if any(not isinstance(value, int) or not minimum_week <= value <= 60 for value in self.week_numbers):
            result.append(Diagnostic("warning", "invalid_week_numbers", "Номера недель будут восстановлены последовательно."))
        if self.week_data_columns and len(self.week_data_columns) != len(self.resolved_week_columns()):
            result.append(Diagnostic("warning", "week_column_count_mismatch", "Столбцы данных будут выровнены со столбцами недель."))
        if self.week_numbers and len(self.week_numbers) != len(self.resolved_week_columns()):
            result.append(Diagnostic("warning", "week_number_count_mismatch", "Номера недель будут выровнены со столбцами."))
        if self.day_start_rows and len(self.day_start_rows) != len(self.day_names):
            result.append(Diagnostic("warning", "day_row_count_mismatch", "Строки дней будут восстановлены по высоте блока."))
        if worksheet:
            if self.weeks_row > worksheet.max_row or self.grid_start_row > worksheet.max_row:
                result.append(Diagnostic("warning", "row_out_of_range", "Координаты будут приведены к границам листа."))
            for col in self.resolved_week_columns():
                if col > worksheet.max_column:
                    result.append(Diagnostic("warning", "column_out_of_range", f"Столбец {col} будет исключён из разметки."))
        return result


from .worksheet_matrix import WorksheetMatrix  # compatibility re-export
