"""Structural fingerprints used to shortlist and learn Excel layout templates."""
from __future__ import annotations

from collections import Counter
import hashlib
import json
from pathlib import Path
import re
from typing import Any, Dict, List

from openpyxl import load_workbook

from .schedule_analyzer import normalize_text, value_as_int

KEYWORDS = (
    "недел", "дисцип", "лектор", "преподав", "обозн", "ауд", "месяц",
    "понедель", "вторник", "сред", "четвер", "пятниц", "суббот",
)


def _hash(payload: Any) -> str:
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _sheet_fingerprint(sheet: Any) -> Dict[str, Any]:
    max_row = int(sheet.max_row or 0)
    max_column = int(sheet.max_column or 0)
    nonempty = 0
    keywords = Counter()
    header_tokens: List[str] = []
    week_positions: List[List[int]] = []
    scan_rows = min(max_row, 140)
    scan_columns = min(max_column, 90)

    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=scan_rows, min_col=1, max_col=scan_columns),
        1,
    ):
        row_weeks: List[int] = []
        for column_index, cell in enumerate(row, 1):
            text = normalize_text(cell.value)
            if not text:
                continue
            nonempty += 1
            lowered = text.casefold().replace("ё", "е")
            for keyword in KEYWORDS:
                if keyword in lowered:
                    keywords[keyword] += 1
            number = value_as_int(cell.value)
            if number is not None and 1 <= number <= 60:
                row_weeks.append(column_index)
            if row_index <= 24 and len(text) <= 120:
                compact = re.sub(r"\s+", " ", lowered).strip()
                if compact and compact not in header_tokens:
                    header_tokens.append(compact)
        if len(row_weeks) >= 3:
            week_positions.append([row_index, min(row_weeks), max(row_weeks), len(row_weeks)])

    header_tokens = header_tokens[:64]
    payload = {
        "title": str(sheet.title),
        "max_row": max_row,
        "max_column": max_column,
        "merged_count": len(sheet.merged_cells.ranges),
        "nonempty_count": nonempty,
        "keyword_count": sum(keywords.values()),
        "keywords": dict(keywords),
        "week_count": max((item[3] for item in week_positions), default=0),
        "week_positions": week_positions[:12],
        "header_tokens": header_tokens,
    }
    payload["signature"] = _hash(payload)
    return payload


def build_workbook_fingerprint(path: str | Path) -> Dict[str, Any]:
    workbook = load_workbook(Path(path), data_only=True, read_only=False)
    sheets = [_sheet_fingerprint(sheet) for sheet in workbook.worksheets]
    payload = {
        "sheet_count": len(sheets),
        "sheets": sheets,
    }
    payload["signature"] = _hash(payload)
    return payload


def selected_sheet_fingerprint(
    fingerprint: Dict[str, Any] | None,
    sheet_name: str,
) -> Dict[str, Any]:
    if not isinstance(fingerprint, dict):
        return {}
    for sheet in fingerprint.get("sheets", []) or []:
        if isinstance(sheet, dict) and str(sheet.get("title")) == str(sheet_name):
            return dict(sheet)
    return {}
