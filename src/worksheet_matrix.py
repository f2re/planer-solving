"""Merged-cell-aware worksheet access without mutating source workbooks."""
from __future__ import annotations

from typing import Any, Dict, List, Optional, Tuple

from openpyxl.worksheet.worksheet import Worksheet


def _present(value: Any) -> bool:
    return value is not None and bool(str(value).replace("\xa0", " ").strip())


class WorksheetMatrix:
    """Merged-cell-aware access without mutating the workbook."""

    def __init__(self, worksheet: Worksheet):
        self.worksheet = worksheet
        self._anchors: Dict[Tuple[int, int], Tuple[int, int]] = {}
        self._ranges: Dict[Tuple[int, int], Tuple[int, int, int, int]] = {}
        for area in worksheet.merged_cells.ranges:
            anchor = (area.min_row, area.min_col)
            bounds = (area.min_row, area.min_col, area.max_row, area.max_col)
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    self._anchors[(row, col)] = anchor
                    self._ranges[(row, col)] = bounds

    def value(self, row: int, col: Optional[int]) -> Any:
        if col is None or row < 1 or col < 1:
            return None
        anchor = self._anchors.get((row, col), (row, col))
        return self.worksheet.cell(*anchor).value

    def is_merged(self, row: int, col: int) -> bool:
        return (row, col) in self._anchors

    def is_anchor(self, row: int, col: int) -> bool:
        return self._anchors.get((row, col), (row, col)) == (row, col)

    def row_value(self, row: int, col: Optional[int]) -> Any:
        """Merged value only when its anchor belongs to the same row.

        Horizontal merges are preserved, while a multi-row header is not
        repeated into the rows below and mistaken for data.
        """
        if col is None or row < 1 or col < 1:
            return None
        anchor_row, anchor_col = self._anchors.get((row, col), (row, col))
        if anchor_row != row:
            return None
        return self.worksheet.cell(anchor_row, anchor_col).value

    def merged_anchor(self, row: int, col: int) -> Tuple[int, int]:
        return self._anchors.get((row, col), (row, col))

    def merged_bounds(self, row: int, col: int) -> Tuple[int, int, int, int]:
        return self._ranges.get((row, col), (row, col, row, col))

    def anchor_values(self, row: int, col_start: int = 1, col_end: Optional[int] = None) -> List[Tuple[int, Any]]:
        end = col_end or self.worksheet.max_column
        return [
            (col, self.value(row, col))
            for col in range(max(1, col_start), min(end, self.worksheet.max_column) + 1)
            if self.is_anchor(row, col) and _present(self.value(row, col))
        ]
