import json
from pathlib import Path

from openpyxl import Workbook

from src.data_loader import DataLoader
from src.schedule_analyzer import ScheduleAnalyzer, ScheduleLayout


def write_teachers(path: Path) -> Path:
    path.write_text(json.dumps([{
        "id": 1,
        "short_name": "Иванов И.И.",
        "full_name": "Иванов Иван Иванович",
        "position": "доцент",
        "rank": "",
        "academic_degree": "",
    }], ensure_ascii=False), encoding="utf-8")
    return path


def build_layered_schedule(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Группа 522"
    sheet["A2"] = "Весенний семестр"
    sheet["A3"] = "2026/2027 учебный год"
    sheet["A7"] = "Уч. недели"
    for column, week in zip(range(5, 8), range(1, 4)):
        sheet.cell(7, column, week)
        sheet.cell(8, column, "Февраль")
    for day_index, day_name in enumerate(["Пн", "Вт", "Ср", "Чт", "Пт", "Сб"]):
        day_start = 10 + day_index * 13
        sheet.cell(day_start, 1, day_name)
        for column in range(5, 8):
            sheet.cell(day_start - 1, column, 10 + day_index)
            for pair_index in range(4):
                pair_row = day_start + pair_index * 3
                sheet.cell(pair_row, column, "Л/Т.01" if pair_index == 0 else "П/Т.01")
                sheet.cell(pair_row + 1, column, "MET")
                sheet.cell(pair_row + 2, column, str(100 + pair_index))
    sheet["A90"] = "Обозн."
    sheet["B90"] = "Наименование дисциплины"
    sheet["J90"] = "Лектор"
    sheet["N90"] = "Другие преподаватели"
    sheet["A93"] = "MET"
    sheet["B93"] = "Метеорология"
    sheet["J93"] = "доц. Иванов И.И."
    sheet["N93"] = "Иванов И.И."
    workbook.save(path)


def test_shifted_grid_and_legend_are_detected(tmp_path: Path) -> None:
    schedule = tmp_path / "522.xlsx"
    build_layered_schedule(schedule)
    layout = ScheduleAnalyzer().analyze(str(schedule)).layout
    assert (layout.weeks_row, layout.first_week_col, layout.last_week_col) == (7, 5, 7)
    assert (layout.grid_start_row, layout.pair_row_stride, layout.pairs_per_day) == (10, 3, 4)
    assert (layout.legend_start_row, layout.legend_code_col) == (90, 1)
    assert (layout.legend_subject_col, layout.legend_lecturer_col, layout.legend_other_col) == (2, 10, 14)


def test_confirmed_layout_parses_lessons(tmp_path: Path) -> None:
    schedule = tmp_path / "522.xlsx"
    teachers = write_teachers(tmp_path / "teachers.json")
    build_layered_schedule(schedule)
    layout = ScheduleAnalyzer().analyze(str(schedule)).layout
    loader = DataLoader(str(teachers))
    lessons = loader.load_group_schedule(str(schedule), group_name="522", layout=layout)
    assert len(lessons) == 72
    assert lessons[0].subject == "MET"
    assert lessons[0].teacher == "Иванов И.И."
    assert loader.last_report["unknown_teacher_lessons"] == 0


def test_combined_cell_uses_embedded_teacher(tmp_path: Path) -> None:
    schedule = tmp_path / "combined.xlsx"
    teachers = write_teachers(tmp_path / "teachers.json")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Нестандартный лист"
    sheet["H12"], sheet["I12"] = 1, 2
    sheet["H13"], sheet["I13"] = "Март", "Март"
    sheet["H15"], sheet["I15"] = "02.03", "09.03"
    sheet["H16"] = sheet["I16"] = "Л/Т.01\nMET\n201\nИванов И.И."
    workbook.save(schedule)
    layout = ScheduleLayout(
        sheet_name="Нестандартный лист", weeks_row=12, first_week_col=8,
        last_week_col=9, months_row=13, grid_start_row=16, grid_end_row=16,
        day_block_rows=1, pairs_per_day=1, pair_row_stride=1, date_row_offset=-1,
        day_names=["Пн"], legend_start_row=None, legend_code_col=None,
        cell_mode="combined_cell", teacher_source="schedule",
    )
    lessons = DataLoader(str(teachers)).load_group_schedule(str(schedule), "Н-1", layout)
    assert len(lessons) == 2
    assert (lessons[0].subject, lessons[0].room, lessons[0].teacher, lessons[0].date_day) == (
        "MET", "201", "Иванов И.И.", 2
    )
