import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from web.backend.app_factory import create_app


def test_all_unreadable_files_still_produce_a_recovery_workbook(tmp_path):
    app = create_app(tmp_path)
    client = TestClient(app)
    workspace_id = client.get("/api/workspaces").json()[0]["id"]
    session_id, session_dir = app.state.context.sessions.create()
    file_id = str(uuid.uuid4())
    stored_name = f"{file_id}.xlsx"
    stored = session_dir / stored_name
    stored.write_bytes(b"not an xlsx archive")
    app.state.context.sessions.save_manifest(
        session_id,
        {
            "session_id": session_id,
            "created_at": 1.0,
            "files": [{
                "file_id": file_id,
                "filename": "broken.xlsx",
                "group_name": "BROKEN",
                "stored_name": stored_name,
                "status": "error",
                "message": "Книгу не удалось разобрать.",
                "analysis": None,
                "bytes_written": stored.stat().st_size,
            }],
        },
    )

    response = client.post(
        f"/api/analysis/{session_id}/generate",
        json={
            "workspace_id": workspace_id,
            "allow_partial": True,
            "files": [{
                "file_id": file_id,
                "group_name": "BROKEN",
                "enabled": True,
                "layout": {},
                "period_overrides": {},
            }],
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "warning"
    assert payload["filename"].startswith("schedule_recovery_")
    assert payload["details"][0]["file_id"] == file_id
    assert payload["details"][0]["used"] is False
    assert payload["details"][0]["status"] == "warning"
    assert payload["details"][0]["action_count"] == 1
    assert payload["reports"][0]["status"] == "skipped"
    assert payload["warnings"]

    output = tmp_path / "output" / payload["filename"]
    assert output.is_file()
    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["Результат разбора", "Замечания и решения"]
    rows = list(workbook["Результат разбора"].iter_rows(values_only=True))
    assert any("broken.xlsx" in row for row in rows)
    assert any("Повторить разбор этого файла" in str(row) for row in rows)
