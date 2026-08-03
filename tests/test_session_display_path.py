import os
import uuid

from openpyxl import Workbook, load_workbook
import pytest

from web.backend.app_context import AnalysisSessionStore
from web.backend.errors import UploadedFileNotFound


def _write_workbook(path, subject):
    workbook = Workbook()
    workbook.active["A1"] = subject
    workbook.save(path)


def test_stored_path_uses_original_filename_for_legacy_uuid_payload(tmp_path):
    store = AnalysisSessionStore(tmp_path / "sessions")
    session_id, session_dir = store.create()
    file_id = str(uuid.uuid4())
    internal = session_dir / f"{file_id}.xlsx"
    _write_workbook(internal, "Метеорология")
    item = {
        "file_id": file_id,
        "filename": "Расписание группы 101.xlsx",
        "stored_name": internal.name,
    }

    display = store.stored_path(session_id, item)
    assert display.name == "Расписание группы 101.xlsx"
    assert display.is_file()
    assert load_workbook(display, data_only=True).active["A1"].value == "Метеорология"
    assert internal.is_file()
    assert session_dir in display.resolve().parents


def test_display_alias_tracks_atomic_replacement_of_canonical_file(tmp_path):
    store = AnalysisSessionStore(tmp_path / "sessions")
    session_id, session_dir = store.create()
    file_id = str(uuid.uuid4())
    internal = session_dir / f"{file_id}.xlsx"
    _write_workbook(internal, "Старая версия")
    item = {
        "file_id": file_id,
        "filename": "Исходное расписание.xlsx",
        "stored_name": internal.name,
    }
    first = store.stored_path(session_id, item)
    assert load_workbook(first, data_only=True).active["A1"].value == "Старая версия"

    replacement = session_dir / ".new.xlsx"
    _write_workbook(replacement, "Новая версия")
    os.replace(replacement, internal)

    second = store.stored_path(session_id, item)
    assert second.name == "Исходное расписание.xlsx"
    assert load_workbook(second, data_only=True).active["A1"].value == "Новая версия"


def test_stored_path_accepts_nested_internal_names_but_rejects_escape(tmp_path):
    store = AnalysisSessionStore(tmp_path / "sessions")
    session_id, session_dir = store.create()
    nested = session_dir / "sources" / "opaque.xlsx"
    nested.parent.mkdir()
    _write_workbook(nested, "Вложенный файл")
    item = {
        "file_id": str(uuid.uuid4()),
        "filename": "Понятное имя.xlsx",
        "stored_name": "sources/opaque.xlsx",
    }
    assert store.stored_path(session_id, item).name == "Понятное имя.xlsx"

    outside = tmp_path / "outside.xlsx"
    _write_workbook(outside, "Нельзя")
    with pytest.raises(UploadedFileNotFound):
        store.stored_path(
            session_id,
            {"file_id": str(uuid.uuid4()), "filename": "x.xlsx", "stored_name": "../../outside.xlsx"},
        )
