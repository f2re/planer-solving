import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.data_loader import DataLoader
from src.schedule_analyzer import ScheduleLayout
from src.transformer import transform_to_teacher_grid
from src.weekly_exporter import generate_weekly_semester_schedule


DAY_NAMES = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб"]


def teachers_file(tmp_path: Path) -> Path:
    path = tmp_path / "teachers.json"
    path.write_text(json.dumps([{
        "id": 1,
        "short_name": "Иванов И.И.",
        "full_name": "Иванов Иван Иванович",
    }], ensure_ascii=False), encoding="utf-8")
    return path


def schedule_file(tmp_path: Path, name: str, *, semester: str, year: str, months: list[str], dates: list[list[int]]) -> tuple[Path, ScheduleLayout]:
    path = tmp_path / name
    wb = Workbook()
    ws = wb.active
    ws.title = "Расписание"
    ws["A1"] = semester
    ws["A2"] = f"{year} учебный год"
    week_columns = [4 + index for index in range(len(months))]
    week_numbers = list(range(1, len(months) + 1))
    day_rows = [8, 13, 18, 23, 28, 33]
    ws.cell(5, 1, "Уч. недели")
    for week, col, month in zip(week_numbers, week_columns, months):
        ws.cell(5, col, week)
        ws.cell(6, col, month)
    for day_index, day_row in enumerate(day_rows):
        ws.cell(day_row, 1, DAY_NAMES[day_index])
        for week_index, col in enumerate(week_columns):
            ws.cell(day_row - 1, col, dates[week_index][day_index])
            ws.cell(day_row, col, "Л")
            ws.cell(day_row + 1, col, "Метеорология")
            ws.cell(day_row + 2, col, "101")
            ws.cell(day_row + 3, col, "Иванов И.И.")
    wb.save(path)
    return path, ScheduleLayout(
        sheet_name="Расписание",
        weeks_row=5,
        first_week_col=week_columns[0],
        last_week_col=week_columns[-1],
        week_columns=week_columns,
        week_data_columns=week_columns,
        week_numbers=week_numbers,
        months_row=6,
        grid_start_row=day_rows[0],
        grid_end_row=36,
        day_start_rows=day_rows,
        pair_row_offsets=[0],
        pairs_per_day=1,
        teacher_row_offset=3,
        teacher_source="schedule",
    )


def test_parser_preserves_month_transition_and_detects_second_semester(tmp_path: Path):
    teachers = teachers_file(tmp_path)
    spring, spring_layout = schedule_file(
        tmp_path,
        "spring.xlsx",
        semester="Весенний семестр",
        year="2025/2026",
        months=["Февраль", "Март"],
        dates=[[23, 24, 25, 26, 27, 28], [2, 3, 4, 5, 6, 7]],
    )
    autumn, autumn_layout = schedule_file(
        tmp_path,
        "autumn.xlsx",
        semester="Осенний семестр",
        year="2026/2027",
        months=["Сентябрь"],
        dates=[[7, 8, 9, 10, 11, 12]],
    )

    loader = DataLoader(str(teachers))
    lessons = loader.load_group_schedule(str(spring), "101", spring_layout)
    assert lessons
    assert next(item for item in lessons if item.week == 1 and item.day_of_week == "Сб").month == "Февраль"
    assert next(item for item in lessons if item.week == 2 and item.day_of_week == "Пн").month == "Март"
    assert loader.last_report["period"]["months"] == ["Февраль", "Март"]

    wrong = loader.load_group_schedule(str(autumn), "102", autumn_layout)
    assert wrong == []
    assert any("разным семестрам" in item for item in loader.last_report["errors"])


def test_parser_rejects_header_and_month_row_conflict(tmp_path: Path):
    teachers = teachers_file(tmp_path)
    path, layout = schedule_file(
        tmp_path,
        "wrong-header.xlsx",
        semester="Весенний семестр",
        year="2025/2026",
        months=["Сентябрь"],
        dates=[[7, 8, 9, 10, 11, 12]],
    )
    loader = DataLoader(str(teachers))
    assert loader.load_group_schedule(str(path), "101", layout) == []
    assert any("Заголовок" in item for item in loader.last_report["errors"])


def test_summary_and_weekly_export_share_the_exact_source_calendar(tmp_path: Path):
    teachers_path = teachers_file(tmp_path)
    teachers = json.loads(teachers_path.read_text(encoding="utf-8"))
    path, layout = schedule_file(
        tmp_path,
        "spring.xlsx",
        semester="Весенний семестр",
        year="2025/2026",
        months=["Февраль", "Март"],
        dates=[[23, 24, 25, 26, 27, 28], [2, 3, 4, 5, 6, 7]],
    )
    loader = DataLoader(str(teachers_path))
    lessons = loader.load_group_schedule(str(path), "101", layout)
    transformed = transform_to_teacher_grid(
        lessons,
        teachers,
        start_date_str="2026-02-01",
        end_date_str="2026-06-30",
        period_reports=[loader.last_report],
    )
    assert transformed["dates"][0] == ("Февраль", 23, 1)
    assert transformed["dates"][6] == ("Март", 2, 2)

    template = tmp_path / "Недельное.xlsx"
    wb = Workbook()
    ws = wb.active
    for row in range(1, 15):
        for col in range(1, 11):
            ws.cell(row, col, "")
    wb.save(template)
    output = tmp_path / "weekly.xlsx"
    generate_weekly_semester_schedule(
        teachers_config=teachers,
        lessons=lessons,
        template_path=str(template),
        output_path=str(output),
        start_date_str="2026-02-01",
        end_date_str="2026-06-30",
        period_reports=[loader.last_report],
        resolved_calendar=transformed["_calendar"],
    )
    result = load_workbook(output, data_only=True)
    assert "23.02" in str(result["Неделя 1"].cell(11, 5).value)
    assert "02.03" in str(result["Неделя 2"].cell(11, 5).value)


def test_parser_rolls_month_inside_one_week(tmp_path: Path):
    teachers = teachers_file(tmp_path)
    path, layout = schedule_file(
        tmp_path,
        "month-boundary.xlsx",
        semester="Весенний семестр",
        year="2025/2026",
        months=["Февраль"],
        dates=[[23, 24, 25, 26, 27, 28]],
    )
    # Replace the dates with the real leap-year week 26 February–2 March.
    wb = load_workbook(path)
    ws = wb.active
    ws["A2"] = "2023/2024 учебный год"
    for row, value in zip([7, 12, 17, 22, 27, 32], [26, 27, 28, 29, 1, 2]):
        ws.cell(row, 4, value)
    wb.save(path)

    loader = DataLoader(str(teachers))
    lessons = loader.load_group_schedule(str(path), "101", layout)
    assert lessons
    assert next(item for item in lessons if item.day_of_week == "Чт").month == "Февраль"
    assert next(item for item in lessons if item.day_of_week == "Пт").month == "Март"
    assert loader.last_report["period"]["months"] == ["Февраль", "Март"]
