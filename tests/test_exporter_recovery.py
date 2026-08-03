from pathlib import Path

from openpyxl import load_workbook

import src.exporter as exporter


def transformed() -> dict:
    return {
        "dates": [("Февраль", 9, 1)],
        "grid": {
            ("Иванов И.И.", 1, "Февраль", 9): {
                "groups": ["522"],
                "subject": "Метеорология",
                "type": "Л",
                "room": "101",
            }
        },
        "grid_vertical": {
            ("Иванов И.И.", 1, "Пн", 1): {
                "groups": ["522"],
                "subject": "Метеорология",
                "type": "Л",
                "room": "101",
            }
        },
        "weeks": [1],
        "week_to_month": {1: "Февраль"},
        "week_day_to_date": {(1, "Пн"): ("Февраль", 9)},
        "semester_info": "Весенний семестр",
        "year_info": "2025/2026",
        "period_report": {"blocking": False},
    }


def test_exporter_sanitizes_duplicate_and_empty_teacher_sheet_names(tmp_path: Path) -> None:
    destination = tmp_path / "schedule.xlsx"
    teachers = [
        {"short_name": "Иванов И.И.", "full_name": "Иванов Иван Иванович"},
        {"short_name": "Иванов И.И.", "full_name": "Иванов Иван Иванович"},
        {"short_name": "", "full_name": ""},
    ]
    result = exporter.export_to_excel(transformed(), teachers, str(destination))
    assert result["mode"] == "normal"
    workbook = load_workbook(destination, data_only=True)
    assert len(workbook.sheetnames) == 4
    assert len(set(workbook.sheetnames)) == 4
    assert all(len(name) <= 31 for name in workbook.sheetnames)


def test_exporter_writes_flat_result_when_normal_formatting_crashes(tmp_path: Path, monkeypatch) -> None:
    destination = tmp_path / "fallback.xlsx"

    def fail(*_args, **_kwargs):
        raise ValueError("synthetic formatting failure")

    monkeypatch.setattr(exporter, "_apply_summary_header", fail)
    result = exporter.export_to_excel(
        transformed(),
        [{"short_name": "Иванов И.И.", "full_name": "Иванов Иван Иванович"}],
        str(destination),
    )
    assert result["mode"] == "fallback"
    assert "synthetic formatting failure" in result["formatting_error"]
    workbook = load_workbook(destination, data_only=True)
    assert workbook.sheetnames == ["Восстановленный результат", "Состав результата"]
    rows = list(workbook["Восстановленный результат"].iter_rows(values_only=True))
    assert any("Метеорология" in row for row in rows)
    assert any("Основное оформление не создано" in str(row) for row in rows)
