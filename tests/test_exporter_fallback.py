from pathlib import Path

from openpyxl import load_workbook

import src.exporter as exporter


def test_presentation_failure_still_creates_downloadable_workbook(
    tmp_path: Path,
    monkeypatch,
) -> None:
    def fail_header(*_args, **_kwargs):
        raise RuntimeError("synthetic formatting failure")

    monkeypatch.setattr(exporter, "_apply_summary_header", fail_header)
    output = tmp_path / "schedule.xlsx"
    transformed = {
        "dates": [["Январь", 15, 1]],
        "grid": {
            ("Иванов И.И.", 1, "Январь", 15): {
                "groups": ["101"],
                "type": "Л",
                "subject": "Математика",
                "room": "201",
            }
        },
        "period_report": {"status": "resolved"},
    }
    teachers = [{
        "id": 1,
        "short_name": "Иванов И.И.",
        "full_name": "Иванов Иван Иванович",
        "position": "доцент",
        "rank": "кандидат наук",
    }]

    result = exporter.export_to_excel(transformed, teachers, str(output))

    assert result["mode"] == "fallback"
    assert "synthetic formatting failure" in result["formatting_error"]
    assert output.is_file()
    workbook = load_workbook(output, data_only=True)
    assert "Восстановленный результат" in workbook.sheetnames
    sheet = workbook["Восстановленный результат"]
    assert sheet["B1"].value == "Основное оформление не создано; данные сохранены в плоской таблице."
    assert any(
        cell.value == "Математика"
        for row in sheet.iter_rows()
        for cell in row
    )
