from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from web.backend.app_factory import create_app


def build_large_schedule(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Большой лист"
    sheet.cell(5, 1, "Уч. недели")
    for index, column in enumerate(range(4, 74), 1):
        sheet.cell(5, column, index)
        sheet.cell(6, column, "Февраль")
    day_rows = [8, 21, 34, 47, 60, 73]
    for day, row in zip(("Пн", "Вт", "Ср", "Чт", "Пт", "Сб"), day_rows):
        sheet.cell(row, 1, day)
        sheet.cell(row, 2, "1-2")
        sheet.cell(row, 4, "Л")
        sheet.cell(row + 1, 4, "MET")
        sheet.cell(row + 2, 4, "101")
    sheet.cell(90, 1, "Обозн")
    sheet.cell(90, 2, "Дисциплина")
    sheet.cell(90, 11, "Лектор")
    sheet.cell(90, 15, "Другие виды занятий")
    sheet.cell(93, 1, "MET")
    sheet.cell(93, 2, "Метеорология")
    sheet.cell(93, 11, "Иванов И.И.")
    sheet.cell(220, 100, "край листа")
    workbook.save(path)


def test_full_sheet_preview_extends_normal_region_limits(tmp_path: Path):
    (tmp_path / "web" / "frontend").mkdir(parents=True)
    source = tmp_path / "large.xlsx"
    build_large_schedule(source)
    client = TestClient(create_app(tmp_path))

    with source.open("rb") as stream:
        analyzed = client.post(
            "/api/analyze",
            files={"files": (source.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
    item = analyzed["files"][0]
    response = client.get(
        f"/api/analysis/{analyzed['session_id']}/files/{item['file_id']}/preview",
        params={"region": "custom", "full_sheet": "true", "sheet_name": "Большой лист"},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["full_sheet"] is True
    assert payload["row_end"] >= 220
    assert payload["col_end"] >= 100
    assert payload["truncated_rows"] is False
    assert payload["truncated_columns"] is False
