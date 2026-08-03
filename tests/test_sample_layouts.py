import json
from pathlib import Path

from openpyxl import Workbook

from src.data_loader import DataLoader
from src.schedule_analyzer import ScheduleAnalyzer, ScheduleLayout
from src.transformer import transform_to_teacher_grid


DAY_NAMES = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб"]


def write_teachers(path: Path) -> Path:
    path.write_text(
        json.dumps(
            [
                {
                    "id": 1,
                    "short_name": "Иванов И.И.",
                    "full_name": "Иванов Иван Иванович",
                },
                {
                    "id": 2,
                    "short_name": "Петров П.П.",
                    "full_name": "Петров Петр Петрович",
                },
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def add_grid(
    sheet,
    *,
    weeks_row: int,
    week_columns: list[int],
    week_numbers: list[int],
    day_rows: list[int],
    codes: list[str] | None = None,
    subject: str = "MET",
) -> None:
    sheet.cell(weeks_row, 1, "Уч. недели")
    for col, week in zip(week_columns, week_numbers):
        sheet.cell(weeks_row, col, week)
        sheet.cell(weeks_row + 1, col, "Февраль")
    selected_codes = codes or ["Л/Т.01", "П/Т.01", "П/Т.02", "П/Т.03"]
    for day_index, day_row in enumerate(day_rows):
        sheet.cell(day_row, 1, DAY_NAMES[day_index])
        for pair_index, offset in enumerate((0, 3, 6, 9)):
            base = day_row + offset
            sheet.cell(base, 2, f"{pair_index * 2 + 1}-{pair_index * 2 + 2}")
            for col_index, col in enumerate(week_columns):
                sheet.cell(day_row - 1, col, 2 + day_index + col_index * 7)
                sheet.cell(base, col, selected_codes[pair_index % len(selected_codes)])
                sheet.cell(base + 1, col, subject)
                sheet.cell(base + 2, col, f"{100 + pair_index}")


def add_legend(
    sheet,
    *,
    header_row: int,
    data_row: int,
    code_col: int,
    subject_col: int,
    lecturer_col: int,
    other_col: int,
    lecturer: str = "Иванов И.И.",
    other: str = "Петров П.П.",
    code: str = "MET",
    subject: str = "Метеорология",
) -> None:
    sheet.cell(header_row, code_col, "Обозн")
    sheet.cell(header_row, subject_col, "Дисциплина")
    sheet.cell(header_row, lecturer_col, "Лектор, уч. степень")
    sheet.cell(header_row, other_col, "Другие виды занятий")
    sheet.cell(data_row, code_col, code)
    sheet.cell(data_row, subject_col, subject)
    sheet.cell(data_row, lecturer_col, lecturer)
    sheet.cell(data_row, other_col, other)
    sheet.cell(data_row + 3, code_col, "ОБОЗНАЧЕНИЯ ВИДОВ ЗАНЯТИЙ:")
    sheet.cell(data_row + 4, code_col, "Л - лекция")
    sheet.cell(data_row + 5, code_col, "П - практические занятия")


def test_short_sample_family_keeps_exact_geometry_and_strict_roles(tmp_path: Path) -> None:
    path = tmp_path / "11161.xlsx"
    teachers = write_teachers(tmp_path / "teachers.json")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "11161"
    week_columns = list(range(4, 9))
    day_rows = [8, 21, 34, 47, 60, 73]
    add_grid(
        sheet,
        weeks_row=5,
        week_columns=week_columns,
        week_numbers=list(range(1, 6)),
        day_rows=day_rows,
    )
    add_legend(
        sheet,
        header_row=83,
        data_row=86,
        code_col=1,
        subject_col=2,
        lecturer_col=11,
        other_col=15,
    )
    workbook.save(path)

    analysis = ScheduleAnalyzer().analyze(str(path))
    layout = analysis.layout
    assert layout.week_columns == week_columns
    assert layout.week_numbers == [1, 2, 3, 4, 5]
    assert layout.day_start_rows == day_rows
    assert layout.pair_row_offsets == [0, 3, 6, 9]
    assert (
        layout.legend_start_row,
        layout.legend_data_start_row,
        layout.legend_end_row,
    ) == (83, 86, 86)
    assert (
        layout.legend_code_col,
        layout.legend_subject_col,
        layout.legend_lecturer_col,
        layout.legend_other_col,
    ) == (1, 2, 11, 15)

    loader = DataLoader(str(teachers))
    lessons = loader.load_group_schedule(str(path), "11161", layout)
    assert len(lessons) == 6 * 4 * 5
    assert lessons[0].teacher == "Иванов И.И."
    assert next(item for item in lessons if item.pair_num == 2).teacher == "Петров П.П."
    assert loader.last_report["unknown_teacher_lessons"] == 0


def test_wide_family_ignores_auxiliary_numeric_tables_and_single_codes(tmp_path: Path) -> None:
    path = tmp_path / "542.xlsx"
    teachers = write_teachers(tmp_path / "teachers.json")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "542"
    week_columns = list(range(5, 13))
    day_rows = [10, 23, 36, 49, 62, 75]
    add_grid(
        sheet,
        weeks_row=7,
        week_columns=week_columns,
        week_numbers=list(range(1, 9)),
        day_rows=day_rows,
        codes=["Л", "ПЗ", "Экзамен/Э", "ИКС/ИКС"],
        subject="ФА",
    )
    add_legend(
        sheet,
        header_row=90,
        data_row=93,
        code_col=1,
        subject_col=3,
        lecturer_col=14,
        other_col=21,
        code="ФА",
        subject="Физика атмосферы",
    )
    # A distracting calculation block similar to the real 542 workbook.
    sheet.cell(135, 40, "Расчётная таблица")
    for offset, week in enumerate(range(1, 31), 41):
        sheet.cell(140, offset, week)
    workbook.save(path)

    layout = ScheduleAnalyzer().analyze(str(path)).layout
    assert layout.weeks_row == 7
    assert layout.week_columns == week_columns
    assert layout.day_start_rows == day_rows
    assert layout.legend_start_row == 90
    assert layout.legend_data_start_row == 93
    assert layout.legend_end_row == 93
    assert (layout.legend_code_col, layout.legend_subject_col) == (1, 3)
    assert (layout.legend_lecturer_col, layout.legend_other_col) == (14, 21)

    loader = DataLoader(str(teachers))
    lessons = loader.load_group_schedule(str(path), "542", layout)
    assert len(lessons) == 6 * 4 * 8
    assert {item.lesson_type_code for item in lessons} == {
        "Л", "ПЗ", "Экзамен/Э", "ИКС/ИКС"
    }
    assert loader.last_report["skipped_non_lesson_cells"] == 0


def test_week_zero_is_preserved_in_parser_and_calendar(tmp_path: Path) -> None:
    path = tmp_path / "552.xlsx"
    teachers = write_teachers(tmp_path / "teachers.json")
    workbook = Workbook()
    sheet = workbook.active
    week_columns = list(range(4, 8))
    add_grid(
        sheet,
        weeks_row=5,
        week_columns=week_columns,
        week_numbers=[0, 1, 2, 3],
        day_rows=[8, 21, 34, 47, 60, 73],
    )
    add_legend(
        sheet,
        header_row=83,
        data_row=86,
        code_col=1,
        subject_col=3,
        lecturer_col=14,
        other_col=21,
    )
    workbook.save(path)

    layout = ScheduleAnalyzer().analyze(str(path)).layout
    assert layout.week_numbers == [0, 1, 2, 3]
    assert layout.allow_week_zero is True
    lessons = DataLoader(str(teachers)).load_group_schedule(str(path), "552", layout)
    assert any(item.week == 0 for item in lessons)

    transformed = transform_to_teacher_grid(
        lessons,
        json.loads(teachers.read_text(encoding="utf-8")),
        start_date_str="2026-02-10",
        end_date_str="2026-03-07",
    )
    assert transformed["weeks"][0] == 0
    assert transformed["week_day_to_date"][(0, "Пн")] == ("Февраль", 2)
    assert any(key[2:] == ("Февраль", 2) for key in transformed["grid"])


def test_merged_week_headers_choose_real_data_columns_without_duplicates(tmp_path: Path) -> None:
    path = tmp_path / "merged-weeks.xlsx"
    teachers = write_teachers(tmp_path / "teachers.json")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Merged"
    sheet["A5"] = "Уч. недели"
    sheet.merge_cells("D5:E5")
    sheet["D5"] = 1
    sheet.merge_cells("F5:G5")
    sheet["F5"] = 2
    sheet.merge_cells("D6:E6")
    sheet["D6"] = "Февраль"
    sheet.merge_cells("F6:G6")
    sheet["F6"] = "Февраль"
    sheet["A8"] = "Пн"
    sheet["B8"] = "1-2"
    for col in (5, 7):
        sheet.cell(7, col, 2)
        sheet.cell(8, col, "Л/Т.01")
        sheet.cell(9, col, "MET")
        sheet.cell(10, col, "101")
    add_legend(
        sheet,
        header_row=20,
        data_row=23,
        code_col=1,
        subject_col=2,
        lecturer_col=6,
        other_col=8,
    )
    workbook.save(path)

    layout = ScheduleAnalyzer().analyze(str(path)).layout
    assert layout.week_columns == [4, 6]
    assert layout.week_data_columns == [5, 7]
    assert layout.week_numbers == [1, 2]
    assert layout.day_start_rows == [8]
    assert layout.pair_row_offsets == [0]

    lessons = DataLoader(str(teachers)).load_group_schedule(str(path), "M", layout)
    assert len(lessons) == 2
    assert [item.week for item in lessons] == [1, 2]


def test_teacher_role_fallback_requires_explicit_operator_choice(tmp_path: Path) -> None:
    path = tmp_path / "strict-role.xlsx"
    teachers = write_teachers(tmp_path / "teachers.json")
    workbook = Workbook()
    sheet = workbook.active
    add_grid(
        sheet,
        weeks_row=5,
        week_columns=[4],
        week_numbers=[1],
        day_rows=[8, 21, 34, 47, 60, 73],
        codes=["Л"],
    )
    add_legend(
        sheet,
        header_row=83,
        data_row=86,
        code_col=1,
        subject_col=2,
        lecturer_col=11,
        other_col=15,
        lecturer="",
        other="Петров П.П.",
    )
    workbook.save(path)

    layout = ScheduleAnalyzer().analyze(str(path)).layout
    strict_loader = DataLoader(str(teachers))
    strict_lessons = strict_loader.load_group_schedule(str(path), "S", layout)
    assert strict_lessons[0].teacher == "Unknown"

    permissive = ScheduleLayout.from_dict(layout.to_dict())
    permissive.teacher_role_fallback = "any"
    fallback_loader = DataLoader(str(teachers))
    fallback_lessons = fallback_loader.load_group_schedule(str(path), "S", permissive)
    assert fallback_lessons[0].teacher == "Петров П.П."


def test_weekly_exporter_creates_week_zero_sheet(tmp_path: Path) -> None:
    from src.data_loader import Lesson
    from src.weekly_exporter import generate_weekly_semester_schedule

    template = tmp_path / "Недельное.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 15):
        for col in range(1, 11):
            sheet.cell(row, col, "")
    workbook.save(template)

    output = tmp_path / "weekly.xlsx"
    teachers = [{
        "short_name": "Иванов И.И.",
        "full_name": "Иванов Иван Иванович",
        "rank": "доцент",
    }]
    lessons = [
        Lesson("552", "MET", "Л", "101", 0, "Пн", 1, "Иванов И.И.", 2, "Февраль"),
        Lesson("552", "MET", "Л", "101", 1, "Пн", 1, "Иванов И.И.", 9, "Февраль"),
    ]
    generate_weekly_semester_schedule(
        teachers_config=teachers,
        lessons=lessons,
        template_path=str(template),
        output_path=str(output),
        start_date_str="2026-02-10",
        end_date_str="2026-02-14",
    )
    result = __import__("openpyxl").load_workbook(output, data_only=True)
    assert result.sheetnames[:2] == ["Неделя 0", "Неделя 1"]
    assert "02.02" in str(result["Неделя 0"].cell(11, 5).value)
