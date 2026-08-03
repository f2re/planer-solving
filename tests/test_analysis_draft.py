import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook

from src.schedule_analyzer import ScheduleAnalyzer
from web.backend.app_factory import create_app


def _session_with_workbook(app, tmp_path):
    source = tmp_path / "draft.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расписание"
    sheet["A1"] = "Весенний семестр"
    sheet["A2"] = "2025/2026 учебный год"
    sheet.cell(5, 4, 1)
    sheet.cell(6, 4, "Февраль")
    sheet.cell(7, 4, 9)
    sheet.cell(8, 1, "Пн")
    sheet.cell(8, 4, "Л")
    sheet.cell(9, 4, "Метеорология")
    sheet.cell(10, 4, "101")
    workbook.save(source)

    session_id, session_dir = app.state.context.sessions.create()
    file_id = str(uuid.uuid4())
    stored_name = f"{file_id}.xlsx"
    stored = session_dir / stored_name
    stored.write_bytes(source.read_bytes())
    analysis = ScheduleAnalyzer().analyze(str(stored)).to_dict()
    manifest = {
        "session_id": session_id,
        "created_at": 1.0,
        "files": [{
            "file_id": file_id,
            "filename": source.name,
            "group_name": "draft",
            "stored_name": stored_name,
            "status": "success",
            "message": "ok",
            "analysis": analysis,
            "bytes_written": stored.stat().st_size,
        }],
    }
    app.state.context.sessions.save_manifest(session_id, manifest)
    return session_id, file_id, analysis


def test_analysis_draft_roundtrip_and_app_restart(tmp_path):
    app = create_app(tmp_path)
    client = TestClient(app)
    session_id, file_id, analysis = _session_with_workbook(app, tmp_path)
    payload = {
        "version": 1,
        "workspace_id": "workspace-a",
        "selected_file_id": file_id,
        "step": 2,
        "files": [{"file_id": file_id, "enabled": True, "group_name": "101 МЕТ"}],
        "layouts": {file_id: {**analysis["layout"], "months_row": 7}},
        "period_overrides": {
            file_id: {
                "week_day_dates": {"1:Пн": "2026-02-16"},
                "week_months": {"1": "Февраль"},
            }
        },
        "calendar_overrides": {"policy": "source", "week_day_dates": {}, "week_months": {}},
        "result": None,
    }

    saved = client.put(f"/api/analysis/{session_id}/draft", json=payload)
    assert saved.status_code == 200
    assert saved.json()["bytes"] > 0

    restored = client.get(f"/api/analysis/{session_id}")
    assert restored.status_code == 200
    data = restored.json()
    assert data["draft"]["workspace_id"] == "workspace-a"
    assert data["draft"]["selected_file_id"] == file_id
    assert data["draft"]["layouts"][file_id]["months_row"] == 7
    assert data["draft"]["period_overrides"][file_id]["week_day_dates"]["1:Пн"] == "2026-02-16"
    assert data["files"][0]["group_name"] == "101 МЕТ"
    assert "stored_name" not in data["files"][0]

    # The draft is stored beside the uploaded files, not in process memory.
    restarted = TestClient(create_app(tmp_path))
    after_restart = restarted.get(f"/api/analysis/{session_id}")
    assert after_restart.status_code == 200
    assert after_restart.json()["draft"]["calendar_overrides"]["policy"] == "source"


def test_draft_rejects_unknown_or_duplicate_file_ids(tmp_path):
    app = create_app(tmp_path)
    client = TestClient(app)
    session_id, file_id, _ = _session_with_workbook(app, tmp_path)
    unknown = str(uuid.uuid4())

    response = client.put(
        f"/api/analysis/{session_id}/draft",
        json={
            "workspace_id": "workspace-a",
            "selected_file_id": unknown,
            "files": [{"file_id": unknown, "enabled": True, "group_name": "bad"}],
        },
    )
    assert response.status_code == 409
    assert response.json()["code"] == "draft_file_mismatch"

    duplicate = client.put(
        f"/api/analysis/{session_id}/draft",
        json={
            "selected_file_id": file_id,
            "files": [
                {"file_id": file_id, "enabled": True, "group_name": "one"},
                {"file_id": file_id, "enabled": False, "group_name": "two"},
            ],
        },
    )
    assert duplicate.status_code == 409
    assert duplicate.json()["code"] == "draft_file_duplicate"


def test_delete_session_removes_saved_draft(tmp_path):
    app = create_app(tmp_path)
    client = TestClient(app)
    session_id, file_id, _ = _session_with_workbook(app, tmp_path)
    assert client.put(
        f"/api/analysis/{session_id}/draft",
        json={"selected_file_id": file_id, "files": [{"file_id": file_id}]},
    ).status_code == 200
    assert client.delete(f"/api/analysis/{session_id}").status_code == 200
    assert client.get(f"/api/analysis/{session_id}").status_code == 404
