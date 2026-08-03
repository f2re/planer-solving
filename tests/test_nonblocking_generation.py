import json
from pathlib import Path
import shutil
import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from src.schedule_analyzer import ScheduleAnalyzer
from tests.test_schedule_analyzer import build_layered_schedule
from web.backend.app_factory import create_app


def _teachers(path: Path) -> None:
    path.write_text(json.dumps([{
        "id": 1,
        "short_name": "Иванов И.И.",
        "full_name": "Иванов Иван Иванович",
        "position": "",
        "rank": "",
        "academic_degree": "",
    }], ensure_ascii=False), encoding="utf-8")


def _register_file(app, session_id: str, session_dir: Path, source: Path, group: str) -> dict:
    file_id = str(uuid.uuid4())
    destination = session_dir / f"{file_id}{source.suffix}"
    shutil.copy2(source, destination)
    return {
        "file_id": file_id,
        "filename": source.name,
        "group_name": group,
        "stored_name": destination.name,
        "status": "success",
        "message": "test",
        "analysis": {},
    }


def _empty_schedule(path: Path) -> dict:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расписание"
    sheet["A1"] = "Весенний семестр"
    sheet["A2"] = "2025/2026 учебный год"
    sheet.cell(5, 4, 1)
    sheet.cell(6, 4, "Февраль")
    for row, day, value in zip([8, 13, 18, 23, 28, 33], ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб"], [9, 10, 11, 12, 13, 14]):
        sheet.cell(row, 1, day)
        sheet.cell(row - 1, 4, value)
    workbook.save(path)
    return {
        "sheet_name": "Расписание",
        "weeks_row": 5,
        "first_week_col": 4,
        "last_week_col": 4,
        "week_columns": [4],
        "week_data_columns": [4],
        "week_numbers": [1],
        "months_row": 6,
        "grid_start_row": 8,
        "grid_end_row": 36,
        "day_start_rows": [8, 13, 18, 23, 28, 33],
        "pair_row_offsets": [0],
        "pairs_per_day": 1,
        "teacher_source": "schedule",
        "teacher_row_offset": 3,
    }


def test_zero_lessons_still_returns_diagnostic_workbook(tmp_path: Path) -> None:
    _teachers(tmp_path / "teachers.json")
    source = tmp_path / "empty.xlsx"
    layout = _empty_schedule(source)
    app = create_app(tmp_path)
    client = TestClient(app)
    workspace_id = client.get("/api/workspaces").json()[0]["id"]
    session_id, session_dir = app.state.context.sessions.create()
    item = _register_file(app, session_id, session_dir, source, "101")
    app.state.context.sessions.save_manifest(session_id, {"session_id": session_id, "files": [item]})

    response = client.post(
        f"/api/analysis/{session_id}/generate",
        json={
            "workspace_id": workspace_id,
            "allow_partial": True,
            "files": [{
                "file_id": item["file_id"],
                "group_name": "101",
                "layout": layout,
                "enabled": True,
                "period_overrides": {},
            }],
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "warning"
    assert payload["filename"].startswith("schedule_recovery_")
    assert payload["details"][0]["used"] is False
    assert payload["details"][0]["action_count"] >= 1
    output = tmp_path / "output" / payload["filename"]
    assert output.is_file()
    result = load_workbook(output, data_only=True)
    assert result.sheetnames == ["Результат разбора", "Замечания и решения"]
    assert "повторной загрузки" in str(result["Результат разбора"]["B3"].value)


def test_broken_file_does_not_cancel_valid_files(tmp_path: Path, monkeypatch) -> None:
    _teachers(tmp_path / "teachers.json")
    good = tmp_path / "good.xlsx"
    build_layered_schedule(good)
    good_layout = ScheduleAnalyzer().analyze(str(good)).layout.to_dict()
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not an xlsx archive")

    app = create_app(tmp_path)
    client = TestClient(app)
    workspace_id = client.get("/api/workspaces").json()[0]["id"]
    client.put(
        f"/api/workspaces/{workspace_id}",
        json={"settings": {"schedule_start_date": "2027-02-01", "schedule_end_date": "2027-06-30"}},
    )
    session_id, session_dir = app.state.context.sessions.create()
    good_item = _register_file(app, session_id, session_dir, good, "522")
    broken_item = _register_file(app, session_id, session_dir, broken, "BROKEN")
    app.state.context.sessions.save_manifest(
        session_id,
        {"session_id": session_id, "files": [good_item, broken_item]},
    )

    response = client.post(
        f"/api/analysis/{session_id}/generate",
        json={
            "workspace_id": workspace_id,
            "allow_partial": True,
            "files": [
                {
                    "file_id": good_item["file_id"],
                    "group_name": "522",
                    "layout": good_layout,
                    "enabled": True,
                    "period_overrides": {},
                },
                {
                    "file_id": broken_item["file_id"],
                    "group_name": "BROKEN",
                    "layout": good_layout,
                    "enabled": True,
                    "period_overrides": {},
                },
            ],
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"]
    assert (tmp_path / "output" / payload["filename"]).is_file()
    details = {item["filename"]: item for item in payload["details"]}
    assert details["good.xlsx"]["used"] is True
    assert details["good.xlsx"]["lesson_count"] == 72
    assert details["broken.xlsx"]["used"] is False
    assert details["broken.xlsx"]["status"] == "warning"
    assert payload["status"] == "warning"


def test_validation_accepts_manual_date_without_reupload(tmp_path: Path) -> None:
    _teachers(tmp_path / "teachers.json")
    source = tmp_path / "empty.xlsx"
    layout = _empty_schedule(source)
    app = create_app(tmp_path)
    client = TestClient(app)
    workspace_id = client.get("/api/workspaces").json()[0]["id"]
    session_id, session_dir = app.state.context.sessions.create()
    item = _register_file(app, session_id, session_dir, source, "101")
    app.state.context.sessions.save_manifest(session_id, {"session_id": session_id, "files": [item]})

    response = client.post(
        f"/api/analysis/{session_id}/validate",
        json={
            "file_id": item["file_id"],
            "group_name": "101",
            "layout": layout,
            "workspace_id": workspace_id,
            "period_overrides": {"week_day_dates": {"1:Пн": "2026-02-16"}},
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "warning"
    assert payload["report"]["generation_allowed"] is True
    assert payload["report"]["period"]["week_day_dates"]["1:Пн"]["day"] == 16
    assert payload["report"]["errors"] == []
