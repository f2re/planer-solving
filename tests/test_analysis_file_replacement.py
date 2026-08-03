import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from src.schedule_analyzer import ScheduleAnalyzer
from web.backend.app_factory import create_app


def _workbook(path, subject="Метеорология"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расписание"
    sheet["A1"] = "Весенний семестр"
    sheet["A2"] = "2025/2026 учебный год"
    sheet.cell(5, 4, 1)
    sheet.cell(6, 4, "Февраль")
    sheet.cell(7, 4, 9)
    sheet.cell(8, 1, "Пн")
    sheet.cell(8, 4, "Л")
    sheet.cell(9, 4, subject)
    sheet.cell(10, 4, "101")
    workbook.save(path)


def _session(app, tmp_path):
    session_id, session_dir = app.state.context.sessions.create()
    good_id = str(uuid.uuid4())
    bad_id = str(uuid.uuid4())
    good_source = tmp_path / "good.xlsx"
    _workbook(good_source)
    good_name = f"{good_id}.xlsx"
    good_path = session_dir / good_name
    good_path.write_bytes(good_source.read_bytes())
    good_analysis = ScheduleAnalyzer().analyze(str(good_path)).to_dict()

    bad_name = f"{bad_id}.xlsx"
    bad_path = session_dir / bad_name
    bad_path.write_bytes(b"broken workbook")
    manifest = {
        "session_id": session_id,
        "created_at": 1.0,
        "files": [
            {
                "file_id": good_id,
                "filename": "good.xlsx",
                "group_name": "GOOD",
                "stored_name": good_name,
                "status": "success",
                "message": "ok",
                "analysis": good_analysis,
                "bytes_written": good_path.stat().st_size,
            },
            {
                "file_id": bad_id,
                "filename": "bad.xlsx",
                "group_name": "BAD",
                "stored_name": bad_name,
                "status": "error",
                "message": "broken",
                "analysis": None,
                "bytes_written": bad_path.stat().st_size,
            },
        ],
        "draft": {
            "version": 1,
            "workspace_id": "workspace-a",
            "selected_file_id": good_id,
            "step": 3,
            "files": [
                {"file_id": good_id, "enabled": True, "group_name": "GOOD"},
                {"file_id": bad_id, "enabled": False, "group_name": "BAD"},
            ],
            "layouts": {good_id: good_analysis["layout"], bad_id: {"weeks_row": 99}},
            "period_overrides": {
                good_id: {"week_day_dates": {"1:Пн": "2026-02-16"}},
                bad_id: {"week_months": {"1": "Август"}},
            },
            "calendar_overrides": {"policy": "source"},
            "result": {"filename": "old.xlsx"},
        },
    }
    app.state.context.sessions.save_manifest(session_id, manifest)
    return session_id, good_id, bad_id, bad_path


def test_replace_damaged_file_preserves_other_layouts_and_dates(tmp_path):
    app = create_app(tmp_path)
    client = TestClient(app)
    session_id, good_id, bad_id, old_bad_path = _session(app, tmp_path)
    replacement = tmp_path / "replacement.xlsm"
    _workbook(replacement, subject="Синоптическая метеорология")

    with replacement.open("rb") as stream:
        response = client.post(
            f"/api/analysis/{session_id}/files/{bad_id}/replace",
            files={"files": (replacement.name, stream, "application/vnd.ms-excel.sheet.macroEnabled.12")},
        )
    assert response.status_code == 200
    item = response.json()
    assert item["file_id"] == bad_id
    assert item["analysis"]
    assert item["status"] in {"success", "warning"}
    # The stable internal path is reused so a process interruption cannot leave
    # the manifest pointing at a missing file.
    assert old_bad_path.exists()
    assert load_workbook(old_bad_path, data_only=True).active["D9"].value == "Синоптическая метеорология"

    restored = client.get(f"/api/analysis/{session_id}").json()
    assert len(restored["files"]) == 2
    draft = restored["draft"]
    assert draft["layouts"][good_id]
    assert bad_id not in draft["layouts"]
    assert draft["period_overrides"][good_id]["week_day_dates"]["1:Пн"] == "2026-02-16"
    assert bad_id not in draft["period_overrides"]
    assert draft["calendar_overrides"]["policy"] == "source"
    assert draft["selected_file_id"] == bad_id
    assert draft["step"] == 2
    assert draft["result"] is None
    states = {value["file_id"]: value for value in draft["files"]}
    assert states[bad_id]["enabled"] is True


def test_unreadable_replacement_keeps_previous_file_and_draft(tmp_path):
    app = create_app(tmp_path)
    client = TestClient(app)
    session_id, good_id, bad_id, old_bad_path = _session(app, tmp_path)
    before = old_bad_path.read_bytes()
    unreadable = tmp_path / "still-broken.xlsx"
    unreadable.write_bytes(b"not a zip")

    with unreadable.open("rb") as stream:
        response = client.post(
            f"/api/analysis/{session_id}/files/{bad_id}/replace",
            files={"files": (unreadable.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert response.status_code == 422
    assert response.json()["code"] == "replacement_unreadable"
    assert old_bad_path.read_bytes() == before
    restored = client.get(f"/api/analysis/{session_id}").json()
    assert restored["draft"]["result"]["filename"] == "old.xlsx"
    assert restored["draft"]["layouts"][good_id]


def test_manifest_write_failure_restores_previous_file_and_metadata(tmp_path, monkeypatch):
    app = create_app(tmp_path)
    client = TestClient(app, raise_server_exceptions=False)
    session_id, good_id, bad_id, old_bad_path = _session(app, tmp_path)
    before = old_bad_path.read_bytes()
    original_save = app.state.context.sessions.save_manifest

    def fail_commit(*_args, **_kwargs):
        raise OSError("synthetic manifest failure")

    monkeypatch.setattr(app.state.context.sessions, "save_manifest", fail_commit)
    replacement = tmp_path / "replacement.xlsx"
    _workbook(replacement, subject="Новый предмет")
    with replacement.open("rb") as stream:
        response = client.post(
            f"/api/analysis/{session_id}/files/{bad_id}/replace",
            files={"files": (replacement.name, stream, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert response.status_code == 500
    assert old_bad_path.read_bytes() == before

    monkeypatch.setattr(app.state.context.sessions, "save_manifest", original_save)
    restored = client.get(f"/api/analysis/{session_id}").json()
    bad = next(item for item in restored["files"] if item["file_id"] == bad_id)
    assert bad["filename"] == "bad.xlsx"
    assert bad.get("analysis") is None
    assert restored["draft"]["result"]["filename"] == "old.xlsx"
    assert restored["draft"]["layouts"][good_id]


def test_replacement_rejects_non_excel_file_without_touching_session(tmp_path):
    app = create_app(tmp_path)
    client = TestClient(app)
    session_id, _, bad_id, old_bad_path = _session(app, tmp_path)
    text = tmp_path / "replacement.txt"
    text.write_text("not excel", encoding="utf-8")

    with text.open("rb") as stream:
        response = client.post(
            f"/api/analysis/{session_id}/files/{bad_id}/replace",
            files={"files": (text.name, stream, "text/plain")},
        )
    assert response.status_code == 415
    assert response.json()["code"] == "replacement_file_type"
    assert old_bad_path.exists()
