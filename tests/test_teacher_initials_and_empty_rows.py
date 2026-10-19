import io
import json
from pathlib import Path
import pytest
from openpyxl import Workbook

from src.data_loader import Lesson
from src.import_wizard import classify_teachers, parse_rows, suggest_teacher_mapping
from src.schedule_detection import ScheduleStructureDetector
from src.schedule_layout import ScheduleLayout
from src.schedule_parser_geometry import ScheduleParserGeometry
from src.schedule_analyzer import WorksheetMatrix
from src.sqlite_workspace_store import SQLiteWorkspaceStore
from src.teacher_assignment import resolve_teacher_collisions
from src.teacher_assignment_rules import (
    normalize_rule,
    replace_teacher_assignment_rules,
    list_teacher_assignment_rules,
)
from src.teacher_resolver import (
    TeacherResolver,
    TeacherLookup,
    format_teacher_initials,
    parse_name_identity,
)


def test_parse_name_identity_variations():
    cases = {
        "И.И. Иванов": ("иванов", "и", "и"),
        "И. И. Иванов": ("иванов", "и", "и"),
        "И. Иванов": ("иванов", "и", ""),
        "ИИ Иванов": ("иванов", "и", "и"),
        "И Иванов": ("иванов", "и", ""),
        "Иванов И.И.": ("иванов", "и", "и"),
        "Иванов И. И.": ("иванов", "и", "и"),
        "Иванов И.": ("иванов", "и", ""),
        "Иванов ИИ": ("иванов", "и", "и"),
        "Иванов И": ("иванов", "и", ""),
        "Иванов": ("иванов", "", ""),
        "Иванов Иван Иванович": ("иванов", "и", "и"),
        "Иванов Алексей": ("иванов", "а", ""),
        "Иванов Борис": ("иванов", "б", ""),
        "Иванов А.А.": ("иванов", "а", "а"),
        "Иванов А.Б.": ("иванов", "а", "б"),
        "А. Иванов": ("иванов", "а", ""),
        "АА Иванов": ("иванов", "а", "а"),
        "Мамин-Сибиряк Д.Н.": ("мамин-сибиряк", "д", "н"),
        "Д.Н. Мамин-Сибиряк": ("мамин-сибиряк", "д", "н"),
        "Ли И.И.": ("ли", "и", "и"),
        "И.И. Ли": ("ли", "и", "и"),
    }
    for raw, expected in cases.items():
        assert parse_name_identity(raw) == expected, f"Failed for {raw}"


def test_format_teacher_initials():
    assert format_teacher_initials("Иванов Иван Иванович") == "Иванов И.И."
    assert format_teacher_initials("И.И. Иванов") == "Иванов И.И."
    assert format_teacher_initials("ИИ Иванов") == "Иванов И.И."
    assert format_teacher_initials("И. Иванов") == "Иванов И."
    assert format_teacher_initials("Иванов И") == "Иванов И."
    assert format_teacher_initials("Иванов Алексей") == "Иванов А."
    assert format_teacher_initials("Иванов") == "Иванов"
    assert format_teacher_initials("") == ""


def test_teacher_resolver_disambiguation_different_initials(tmp_path: Path):
    teachers_file = tmp_path / "teachers.json"
    teachers_file.write_text(
        json.dumps([
            {"short_name": "Иванов А.А.", "full_name": "Иванов Алексей Алексеевич"},
            {"short_name": "Иванов Б.Б.", "full_name": "Иванов Борис Борисович"},
            {"short_name": "Петров П.П.", "full_name": "Петров Пётр Петрович"},
        ], ensure_ascii=False),
        encoding="utf-8",
    )
    resolver = TeacherResolver(str(teachers_file))

    # Exact matches
    assert resolver._resolve_teacher_alias("Иванов А.А.") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("А.А. Иванов") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("АА Иванов") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("Иванов АА") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("Иванов Б.Б.") == "Иванов Б.Б."
    assert resolver._resolve_teacher_alias("Б.Б. Иванов") == "Иванов Б.Б."

    # Single initial matching unique teacher
    assert resolver._resolve_teacher_alias("Иванов А.") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("А. Иванов") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("Иванов А") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("А Иванов") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("Иванов Б.") == "Иванов Б.Б."
    assert resolver._resolve_teacher_alias("Б. Иванов") == "Иванов Б.Б."

    # Non-initial surname with multiple matches -> ambiguous
    res = resolver._resolve_teacher_alias("Иванов", warn=True)
    assert res == ""
    assert any("Иванов" in w and "неоднозначен" in w for w in resolver.warnings)

    # Unique surname without initials -> recognized
    assert resolver._resolve_teacher_alias("Петров") == "Петров П.П."
    assert resolver._resolve_teacher_alias("П.П. Петров") == "Петров П.П."
    assert resolver._resolve_teacher_alias("Петров П.") == "Петров П.П."


def test_teacher_resolver_disambiguation_same_first_initial(tmp_path: Path):
    teachers_file = tmp_path / "teachers.json"
    teachers_file.write_text(
        json.dumps([
            {"short_name": "Иванов А.А.", "full_name": "Иванов Алексей Алексеевич"},
            {"short_name": "Иванов А.Б.", "full_name": "Иванов Алексей Борисович"},
        ], ensure_ascii=False),
        encoding="utf-8",
    )
    resolver = TeacherResolver(str(teachers_file))

    # Distinct second initial -> recognized
    assert resolver._resolve_teacher_alias("Иванов А.А.") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("А.А. Иванов") == "Иванов А.А."
    assert resolver._resolve_teacher_alias("Иванов А.Б.") == "Иванов А.Б."
    assert resolver._resolve_teacher_alias("А.Б. Иванов") == "Иванов А.Б."

    # Ambiguous first initial alone
    resolver.warnings.clear()
    assert resolver._resolve_teacher_alias("Иванов А.", warn=True) == ""
    assert any("неоднозначен" in w for w in resolver.warnings)

    resolver.warnings.clear()
    assert resolver._resolve_teacher_alias("А. Иванов", warn=True) == ""
    assert any("неоднозначен" in w for w in resolver.warnings)


def test_teacher_resolver_extract_teachers(tmp_path: Path):
    teachers_file = tmp_path / "teachers.json"
    teachers_file.write_text(
        json.dumps([
            {"short_name": "Иванов А.А.", "full_name": "Иванов Алексей Алексеевич"},
            {"short_name": "Иванов Б.Б.", "full_name": "Иванов Борис Борисович"},
            {"short_name": "Сидоров С.С.", "full_name": "Сидоров Сергей Сергеевич"},
        ], ensure_ascii=False),
        encoding="utf-8",
    )
    resolver = TeacherResolver(str(teachers_file))

    # Single initial in text
    assert resolver._extract_teachers("доц. Иванов А.") == ["Иванов А.А."]
    assert resolver._extract_teachers("доц. А. Иванов") == ["Иванов А.А."]
    assert resolver._extract_teachers("Иванов Б.") == ["Иванов Б.Б."]
    assert resolver._extract_teachers("Иванов АА") == ["Иванов А.А."]
    assert resolver._extract_teachers("АА Иванов") == ["Иванов А.А."]

    # Both teachers mentioned in same cell
    extracted = resolver._extract_teachers("Иванов А.А., Иванов Б.Б.")
    assert set(extracted) == {"Иванов А.А.", "Иванов Б.Б."}

    # Ambiguous bare surname
    resolver.warnings.clear()
    assert resolver._extract_teachers("Иванов") == []
    assert any("Иванов" in w and "неоднозначен" in w for w in resolver.warnings)

    # Single unique surname
    assert resolver._extract_teachers("Сидоров") == ["Сидоров С.С."]
    assert resolver._extract_teachers("С.С. Сидоров") == ["Сидоров С.С."]
    assert resolver._extract_teachers("Сидоров С.") == ["Сидоров С.С."]


def test_teacher_resolver_handles_empty_records_and_none(tmp_path: Path):
    teachers_file = tmp_path / "teachers.json"
    teachers_file.write_text(
        json.dumps([
            None,
            {},
            {"short_name": "", "full_name": ""},
            {"short_name": "Иванов И.И.", "full_name": "Иванов Иван Иванович"},
            None,
            {"short_name": "  "},
        ], ensure_ascii=False),
        encoding="utf-8",
    )
    resolver = TeacherResolver(str(teachers_file))
    assert len(resolver.teachers_config) == 1
    assert resolver._resolve_teacher_alias("Иванов И.И.") == "Иванов И.И."
    assert resolver._resolve_teacher_alias("И.И. Иванов") == "Иванов И.И."
    assert resolver._resolve_teacher_alias("Иванов") == "Иванов И.И."


def test_teacher_assignment_with_initial_variations_and_empty_config():
    teachers = [
        None,
        {},
        {"short_name": "Иванов А.А.", "full_name": "Иванов Алексей Алексеевич"},
        {"short_name": "Иванов Б.Б.", "full_name": "Иванов Борис Борисович"},
        {"short_name": "Сидоров С.С.", "full_name": "Сидоров Сергей Сергеевич"},
    ]
    lessons = [
        Lesson(
            group="G1", subject="Math", lesson_type_code="Л", room="101",
            week=1, day_of_week="Пн", pair_num=1, teacher="А.А. Иванов",
            date_day=1, month="Сентябрь",
        ),
        Lesson(
            group="G2", subject="Physics", lesson_type_code="Л", room="102",
            week=1, day_of_week="Пн", pair_num=1, teacher="Иванов Б.",
            date_day=1, month="Сентябрь",
        ),
        Lesson(
            group="G3", subject="Chemistry", lesson_type_code="П", room="103",
            week=1, day_of_week="Пн", pair_num=1, teacher="Сидоров",
            date_day=1, month="Сентябрь",
        ),
    ]
    report = resolve_teacher_collisions(
        lessons,
        teachers,
    )
    assert lessons[0].teacher == "Иванов А.А."
    assert lessons[1].teacher == "Иванов Б.Б."
    assert lessons[2].teacher == "Сидоров С.С."
    assert report["original_conflict_count"] == 0


def test_teacher_assignment_rules_with_initials():
    teachers = [
        {"short_name": "Иванов И.И.", "full_name": "Иванов Иван Иванович"},
        {"short_name": "Петров П.П.", "full_name": "Петров Пётр Петрович"},
    ]
    lookup = TeacherLookup(teachers)

    rule = normalize_rule(
        {
            "subject": "Математика",
            "lecturer": "И.И. Иванов",
            "practice": "Петров П.",
            "reserve": "ИИ Иванов",
        },
        aliases=lookup,
    )
    assert rule["lecturer"] == "Иванов И.И."
    assert rule["practice"] == "Петров П.П."
    assert rule["reserve"] == "Иванов И.И."


def test_legend_parsing_does_not_break_on_empty_rows(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # Header for legend
    ws.cell(row=5, column=1, value="Код")
    ws.cell(row=5, column=2, value="Дисциплина")
    ws.cell(row=5, column=3, value="Лектор")
    ws.cell(row=5, column=4, value="Другие")

    # Record 1
    ws.cell(row=6, column=1, value="МАТ101")
    ws.cell(row=6, column=2, value="Математика")
    ws.cell(row=6, column=3, value="Иванов И.И.")

    # Two blank rows in between!
    # row 7 is blank
    # row 8 is blank

    # Record 2
    ws.cell(row=9, column=1, value="ФИЗ101")
    ws.cell(row=9, column=2, value="Физика")
    ws.cell(row=9, column=3, value="Петров П.П.")

    matrix = WorksheetMatrix(ws)
    layout = ScheduleLayout(
        sheet_name="Sheet1",
        legend_start_row=5,
        legend_data_start_row=6,
        legend_end_row=9,
        legend_code_col=1,
        legend_subject_col=2,
        legend_lecturer_col=3,
        legend_other_col=4,
    )

    teachers_file = tmp_path / "teachers.json"
    teachers_file.write_text(
        json.dumps([
            {"short_name": "Иванов И.И.", "full_name": "Иванов Иван Иванович"},
            {"short_name": "Петров П.П.", "full_name": "Петров Пётр Петрович"},
        ], ensure_ascii=False),
        encoding="utf-8",
    )

    class DummyParser(ScheduleParserGeometry, TeacherResolver):
        def __init__(self, tpath):
            TeacherResolver.__init__(self, tpath)

    parser = DummyParser(str(teachers_file))
    report = {"warnings": []}
    mapping = parser._legend(matrix, layout, report)

    # Both records should be successfully parsed despite blank rows 7 and 8
    assert "мат101" in mapping or "математика" in mapping
    assert "физ101" in mapping or "физика" in mapping
    assert report["legend_entries"] >= 2


def test_schedule_detection_does_not_break_on_blank_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # Header
    ws.cell(row=1, column=1, value="Шифр")
    ws.cell(row=1, column=2, value="Наименование дисциплины")
    ws.cell(row=1, column=3, value="Лектор")
    ws.cell(row=1, column=4, value="Практика")

    # Record 1
    ws.cell(row=2, column=1, value="DISC1")
    ws.cell(row=2, column=2, value="Математический анализ")
    ws.cell(row=2, column=3, value="Иванов И.И.")

    # 3 blank rows!
    # rows 3, 4, 5 are blank

    # Record 2
    ws.cell(row=6, column=1, value="DISC2")
    ws.cell(row=6, column=2, value="Общая физика")
    ws.cell(row=6, column=3, value="Петров П.П.")

    matrix = WorksheetMatrix(ws)
    detector = ScheduleStructureDetector()
    legend_meta = detector._detect_legend(matrix, 1)

    assert legend_meta["start_row"] == 1
    assert legend_meta["data_start_row"] == 2
    # End row should reach row 6 despite blanks on rows 3-5
    assert legend_meta["end_row"] >= 6


def test_import_wizard_safely_ignores_empty_rows():
    csv_payload = (
        "ФИО;Краткое имя;Должность\n"
        "Иванов Иван Иванович;Иванов И.И.;Доцент\n"
        ";;\n"
        "   ;   ;   \n"
        "Петров Петр Петрович;Петров П.П.;Профессор\n"
        ";;\n"
    ).encode("utf-8")

    source = parse_rows(csv_payload, "teachers.csv")
    mapping = suggest_teacher_mapping(source["headers"])
    # Pass extra empty row in rows to classify_teachers
    rows_with_empties = list(source["rows"]) + [{}, {"ФИО": "", "Должность": ""}]
    preview = classify_teachers(rows_with_empties, mapping, [])

    # The empty rows should be safely skipped, not counted as invalid
    assert preview["counts"]["invalid"] == 0
    assert preview["counts"]["add"] == 2
    assert len(preview["rows"]) == 2
    names = [r["teacher"]["short_name"] for r in preview["rows"]]
    assert "Иванов И.И." in names
    assert "Петров П.П." in names


def test_sqlite_workspace_store_import_teachers_ignores_empty_records(tmp_path: Path):
    store = SQLiteWorkspaceStore(tmp_path / "workspaces.db")
    ws = store.create_workspace({"name": "Test Workspace"})
    workspace_id = ws["id"]

    records = [
        None,
        {},
        {"short_name": "", "full_name": ""},
        {"short_name": "Иванов И.И.", "full_name": "Иванов Иван Иванович", "position": "Доцент"},
        None,
        {"short_name": "  ", "full_name": ""},
        {"short_name": "Петров П.П.", "full_name": "Петров Петр Петрович", "position": "Профессор"},
    ]

    result = store.import_teachers(workspace_id, records)
    assert result["added"] == 2
    assert result["total"] == 2

    teachers = store.list_teachers(workspace_id)
    assert len(teachers) == 2
    short_names = {t["short_name"] for t in teachers}
    assert short_names == {"Иванов И.И.", "Петров П.П."}
